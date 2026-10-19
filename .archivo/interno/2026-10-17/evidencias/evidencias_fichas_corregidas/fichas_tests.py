"""Variantes complementarias B20. Solo base MySQL temporal y correo en memoria."""
import json,re,unittest,time
from pathlib import Path
from datetime import timedelta,datetime
from zoneinfo import ZoneInfo
from io import BytesIO
from unittest.mock import patch
from django.test import TestCase,Client,override_settings
from django.test.runner import DiscoverRunner
from django.core import mail
from django.utils import timezone
from openpyxl import load_workbook
from core.models import Usuario,Rol,Delegacion,Cargo,TipoAtencion,SubAtencion,ItemMeta,ConfiguracionMeta,Atencion,RecuperacionAcceso,AccesoFallido,EventoLocal
from core.local_api import PERMISSIONS
OUT=Path(__file__).resolve().parents[2]/'output/evidencias_fichas_corregidas'
PROOFS=[]
@override_settings(EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
class Fichas(TestCase):
 def setUp(self):
  self.d=Delegacion.objects.create(codigo='D01',nombre='Centro QA',comuna='La Serena')
  self.d2=Delegacion.objects.create(codigo='D02',nombre='Norte QA',comuna='La Serena')
  self.role=Rol.objects.create(nombre='Operativo QA',permisos=['atenciones.ver','atenciones.crear','atenciones.exportar','tipos-atencion.ver','subatenciones.ver'])
  self.cargo=Cargo.objects.create(codigo='CQ',nombre='Cargo QA')
  self.user=Usuario.objects.create_user('ficha.qa',email='ficha@example.test',first_name='Cuenta QA',password='FixtureInicial!347652',rol=self.role,delegacion=self.d,cargo=self.cargo)
  self.user.delegaciones.add(self.d)
  self.admin=Usuario.objects.create_superuser('admin.fichas','admin.fichas@example.test','FixtureAdmin!85762',first_name='Administración QA',delegacion=self.d)
  self.admin.delegaciones.set([self.d,self.d2]);self.client.force_login(self.admin)
  self.tipo=TipoAtencion.objects.create(codigo='T0',nombre='Tipo base')
  self.sub=SubAtencion.objects.create(codigo='S0',nombre='Sub base',tipo_atencion=self.tipo)
 def check(self,case,name,got,want):
  good=got==want;PROOFS.append({'caso':case,'variante':name,'esperado':want,'obtenido':got,'estado':'Aprobado' if good else 'Fallido','fecha':datetime.now(ZoneInfo('America/Santiago')).isoformat(),'metodo':self._testMethodName})
  self.assertEqual(got,want,name)
 def send(self,path,data,method='post',client=None):return getattr(client or self.client,method)(path,json.dumps(data),content_type='application/json')
 def edit(self,url,data):
  current=self.client.get(url).json()['data'];return self.send(url,{'_version':current['version'],**data},'patch')
 def deactivate(self,url):
  current=self.client.get(url).json()['data'];return self.send(url+'toggle-estado/',{'_version':current['version']})
 def payload(self,r,i=1):
  if r=='usuarios':return {'username':f'QAuser{i}','first_name':f'Cuenta {i}','email':f'qa{i}@example.test','password':'FixtureCuenta!975863','rol':self.role.pk,'cargo':self.cargo.pk,'delegacion':self.d.pk,'delegaciones':[self.d.pk]}
  if r=='roles':return {'nombre':f'QA Registro {i}','descripcion':'Prueba controlada','permisos':['atenciones.ver'],'alcance':'delegacion'}
  p={'codigo':f'QA{i}','nombre':f'QA Registro {i}'}
  if r=='delegaciones':p.update(direccion='Dirección de ensayo',comuna='La Serena')
  if r=='items-metas':p.update(descripcion='Ítem de ensayo',unidad_medida='Atención',cargos=[])
  if r=='subatenciones':p['tipo_atencion']=self.tipo.pk
  return p
 def crud(self,r,case):
  base='/api/'+r+'/';p=self.payload(r)
  response=self.send(base,p);self.check(case,'Crear registro válido',response.status_code,201);obj=response.json()['data'];url=base+str(obj['id'])+'/'
  self.check(case,'Consultar registro creado',self.client.get(url).status_code,200)
  statekey='is_active' if r=='usuarios' else 'estado';namekey='first_name' if r=='usuarios' else 'nombre'
  self.check(case,'Editar registro',self.edit(url,{namekey:'QA Editado'}).status_code,200)
  self.check(case,'Persistencia de edición',self.client.get(url).json()['data'][namekey],'QA Editado')
  duplicate=self.payload(r);duplicate[namekey]='QA Editado'
  self.check(case,'Rechazo de duplicado sin guardar',self.send(base,duplicate).status_code,400)
  required=['username','first_name','email'] if r=='usuarios' else (['nombre'] if r=='roles' else ['codigo','nombre'])
  for key in required:
   invalid=self.payload(r,100);invalid[key]=''
   self.check(case,'Campo obligatorio '+key,self.send(base,invalid).status_code,400)
  model={'delegaciones':Delegacion,'roles':Rol,'usuarios':Usuario,'items-metas':ItemMeta,'tipos-atencion':TipoAtencion,'subatenciones':SubAtencion}[r]
  maxfield='username' if r=='usuarios' else 'codigo' if r!='roles' else 'nombre';limit=model._meta.get_field(maxfield).max_length
  boundary=self.payload(r,101);boundary[maxfield]='Z'*limit
  self.check(case,'Longitud máxima '+maxfield,self.send(base,boundary).status_code,201)
  boundary=self.payload(r,102);boundary[maxfield]='Z'*(limit+1)
  self.check(case,'Máximo más uno '+maxfield,self.send(base,boundary).status_code,400)
  before=model.objects.get(pk=obj['id']);v=before.version
  self.check(case,'Sin versión no se sobrescribe',self.send(url,{namekey:'Alterado'},'patch').status_code,428)
  self.check(case,'Versión obsoleta rechazada',self.send(url,{'_version':v-1,namekey:'Alterado'},'patch').status_code,409)
  self.check(case,'Sin mutación tras edición inválida',self.client.get(url).json()['data'][namekey],'QA Editado')
  self.check(case,'Desactivar registro libre',self.deactivate(url).status_code,200)
  self.check(case,'Estado inactivo persistido',self.client.get(url).json()['data'][statekey],False)
  self.check(case,'Reactivar registro libre',self.deactivate(url).status_code,200)
  self.check(case,'Estado activo persistido',self.client.get(url).json()['data'][statekey],True)
  # 23 coincidencias exportables y dos no coincidentes, sin usar datos personales reales.
  ids=[]
  for i in range(200,225):
   value=self.payload(r,i)
   key='first_name' if r=='usuarios' else 'nombre';value[key]=('FiltroFicha ' if i<223 else 'FueraFicha ')+str(i)
   saved=self.send(base,value);self.assertEqual(saved.status_code,201);ids.append(saved.json()['data']['id'])
  visible=[]
  for page in [1,2,3]:
   body=self.client.get(base,{'search':'filtroficha','page_size':10,'page':page}).json()
   self.check(case,'Total filtrado página '+str(page),body['count'],23);visible.extend(o['id']for o in body['results'])
  self.check(case,'Paginación sin omisiones ni duplicados',sorted(visible),sorted(ids[:23]))
  x=self.client.get(base+'export/',{'search':'FILTROFICHA','page_size':10,'page':3});sheet=load_workbook(BytesIO(x.content)).active
  exportids=[row[0]for row in sheet.iter_rows(min_row=2,values_only=True)]
  self.check(case,'XLSX exporta las 23 coincidencias de todas las páginas',sorted(exportids),sorted(ids[:23]))
  headers=[str(c.value)for c in sheet[1]]
  self.check(case,'XLSX sin columnas de credenciales',any(any(secret in h.lower()for secret in ['password','token','hash','codigo_recuperacion'])for h in headers),False)
  (OUT/f'{r}-23-resultados.xlsx').write_bytes(x.content)
  self.check(case,'Sin coincidencias explícito',self.client.get(base,{'search':'INEXISTENTE_FICHA'}).json()['count'],0)
  self.check(case,'Auditoría de edición presente',EventoLocal.objects.filter(recurso=r,registro=obj['id'],accion='editar').exists(),True)
 def test_ADM01_ADM02_delegaciones(self):
  self.crud('delegaciones','ADM-01');url=f'/api/delegaciones/{self.d.pk}/'
  self.check('ADM-02','Baja con usuarios relacionados rechazada',self.deactivate(url).status_code,400)
  self.check('ADM-02','Delegación e historia preservadas',Delegacion.objects.filter(pk=self.d.pk,estado=True).exists(),True)
  self.check('ADM-02','Duplicado de nombre rechazado',self.send('/api/delegaciones/',{'codigo':'NUEVO','nombre':self.d.nombre}).status_code,400)
 def test_ADM03_usuarios(self):
  self.crud('usuarios','ADM-03');url=f'/api/usuarios/{self.user.pk}/'
  other=Rol.objects.create(nombre='Consulta asignada',permisos=['tipos-atencion.ver'])
  cargo=Cargo.objects.create(codigo='OT',nombre='Otro cargo')
  self.check('ADM-03','Editar cargo rol y delegaciones autorizadas',self.edit(url,{'cargo':cargo.pk,'rol':other.pk,'delegacion':self.d2.pk,'delegaciones':[self.d.pk,self.d2.pk]}).status_code,200)
  c=Client();self.check('ADM-03','Login tras asignación',self.send('/api/login/',{'username':self.user.username,'password':'FixtureInicial!347652'},client=c).status_code,200)
  self.check('ADM-03','Operación no autorizada por nuevo rol',self.send('/api/atenciones/',{},client=c).status_code,403)
  self.check('ADM-03','Desactivar cuenta',self.deactivate(url).status_code,200)
  self.check('ADM-03','Cuenta desactivada no entra',self.send('/api/login/',{'username':self.user.username,'password':'FixtureInicial!347652'},client=Client()).status_code,401)
  self.check('ADM-03','Correo inválido rechazado',self.edit(url,{'email':'invalido'}).status_code,400)
  self.check('ADM-03','Delegación no incluida rechazada',self.edit(url,{'delegacion':self.d2.pk,'delegaciones':[self.d.pk]}).status_code,400)
 def test_ADM04_roles(self):
  self.crud('roles','ADM-04');role=Rol.objects.create(nombre='Lectura asignable',permisos=['tipos-atencion.ver'])
  self.check('ADM-04','Asignar rol a cuenta',self.edit(f'/api/usuarios/{self.user.pk}/',{'rol':role.pk}).status_code,200)
  c=Client();c.force_login(self.user)
  self.check('ADM-04','Lectura permitida',c.get('/api/tipos-atencion/').status_code,200)
  self.check('ADM-04','Escritura denegada',self.send('/api/tipos-atencion/',{'codigo':'NO','nombre':'No'},client=c).status_code,403)
  self.check('ADM-04','Baja rol vinculado rechazada',self.deactivate(f'/api/roles/{role.pk}/').status_code,400)
  self.check('ADM-04','Modificar permisos del rol',self.edit(f'/api/roles/{role.pk}/',{'permisos':[]}).status_code,200)
  self.check('ADM-04','Permiso retirado efectivo',c.get('/api/tipos-atencion/').status_code,403)
 def test_CAT01_items(self):
  self.crud('items-metas','CAT-01');item=ItemMeta.objects.create(codigo='REF',nombre='Ítem relacionado');ConfiguracionMeta.objects.create(item=item,periodo='QA',meta_esperada=10,ponderacion=40)
  self.check('CAT-01','Baja de ítem con meta activa rechazada',self.deactivate(f'/api/items-metas/{item.pk}/').status_code,400)
 def test_CAT02_tipos(self):
  self.crud('tipos-atencion','CAT-02')
  self.check('CAT-02','Baja con subatenciones activas rechazada',self.deactivate(f'/api/tipos-atencion/{self.tipo.pk}/').status_code,400)
 def test_CAT03_subatenciones(self):
  self.crud('subatenciones','CAT-03');url=f'/api/subatenciones/{self.sub.pk}/'
  other=TipoAtencion.objects.create(codigo='OT',nombre='Otro padre')
  self.check('CAT-03','Cambiar padre válido',self.edit(url,{'tipo_atencion':other.pk}).status_code,200)
  self.check('CAT-03','Padre inexistente rechazado',self.edit(url,{'tipo_atencion':999999}).status_code,400)
  other.estado=False;other.save()
  self.check('CAT-03','Padre inactivo rechazado',self.edit(url,{'tipo_atencion':other.pk}).status_code,400)
 def test_ACC02_variantes_y_ACC03_logout(self):
  variants=[('Contraseña incorrecta',{'username':self.user.username,'password':'Incorrecta'}),('Usuario inexistente',{'username':'no_existe','password':'Incorrecta'}),('Usuario vacío',{'username':'','password':'Incorrecta'}),('Contraseña vacía',{'username':self.user.username,'password':''}),('Ambos vacíos',{})]
  for label,data in variants:
   AccesoFallido.objects.all().delete();c=Client()
   self.check('ACC-02',label,self.send('/api/login/',data,client=c).status_code,401)
   self.check('ACC-02',label+' sin sesión',c.get('/api/contexto/').status_code,401)
  for label in ['Cuenta inactiva','Rol inactivo']:
   AccesoFallido.objects.all().delete();self.user.is_active=label!='Cuenta inactiva';self.user.save();self.role.estado=label!='Rol inactivo';self.role.save()
   self.check('ACC-02',label,self.send('/api/login/',{'username':self.user.username,'password':'FixtureInicial!347652'},client=Client()).status_code,401)
  self.user.is_active=True;self.user.save();self.role.estado=True;self.role.save();c=Client();c.force_login(self.user)
  self.check('ACC-03','Logout aceptado',self.send('/api/logout/',{},client=c).status_code,200)
  self.check('ACC-03','Petición protegida tras logout',c.get('/api/contexto/').status_code,401)
 def test_ACC04_fallo_de_envio_sin_falsa_entrega(self):
  from smtplib import SMTPException
  with patch('core.local_api.send_mail',side_effect=SMTPException('Fallo controlado de QA')):
   r=self.send('/api/recuperacion/',{'email':self.user.email},client=Client())
  self.check('ACC-04','Fallo de envío conserva respuesta genérica',r.status_code,200)
  self.check('ACC-04','Código no entregado no queda utilizable',RecuperacionAcceso.objects.filter(usuario=self.user,consumido=False).count(),0)
 def test_ACC04_respuesta_generica_y_entrega_controlada(self):
  known=Client();unknown=Client()
  a=self.send('/api/recuperacion/',{'email':self.user.email},client=known);b=self.send('/api/recuperacion/',{'email':'no-registro@example.test'},client=unknown)
  self.check('ACC-04','Respuesta HTTP equivalente',(a.status_code,b.status_code),(200,200))
  self.check('ACC-04','Mensaje exacto equivalente',a.json()['message'],b.json()['message'])
  self.check('ACC-04','Solo cuenta existente recibe mensaje en buzón QA',len(mail.outbox),1)
  self.check('ACC-04','Destinatario de la cuenta correcta',mail.outbox[0].to,[self.user.email])
 def request_code(self,c):
  a=self.send('/api/recuperacion/',{'email':self.user.email},client=c);self.assertEqual(a.status_code,200)
  return re.search(r'\b\d{6}\b',mail.outbox[-1].body).group()
 def test_ACC05_caducidad_intentos_reenvio_y_consumo(self):
  c=Client();code=self.request_code(c)
  self.check('ACC-05','Reenvío antes de 60 s bloqueado por servidor',self.send('/api/recuperacion/',{'email':self.user.email},client=c).status_code,429)
  now=timezone.now()
  with patch('core.local_api.timezone.now',return_value=now+timedelta(seconds=61)):
   self.check('ACC-05','Reenvío después de 60 s permitido',self.send('/api/recuperacion/',{'email':self.user.email},client=c).status_code,200)
  self.check('ACC-05','Segundo mensaje QA generado',len(mail.outbox),2)
  fresh=re.search(r'\b\d{6}\b',mail.outbox[-1].body).group()
  self.check('ACC-05','Solicitud anterior invalidada',RecuperacionAcceso.objects.filter(usuario=self.user,consumido=True).count(),1)
  with patch('core.local_api.timezone.now',return_value=now+timedelta(minutes=12)):
   self.check('ACC-05','Código vencido rechazado',self.send('/api/recuperacion/verificar/',{'codigo':fresh},client=c).status_code,400)
  wrong='000000' if fresh!='000000' else '111111'
  for i in range(5):self.check('ACC-05','Intento incorrecto '+str(i+1),self.send('/api/recuperacion/verificar/',{'codigo':wrong},client=c).status_code,400)
  self.check('ACC-05','Código correcto tras agotar intentos rechazado',self.send('/api/recuperacion/verificar/',{'codigo':fresh},client=c).status_code,400)
 def test_ACC06_cambio_seguro_y_reuso(self):
  c=Client();code=self.request_code(c)
  payload={'password':'FixtureRenovada!648379','confirmacion':'FixtureRenovada!648379'}
  self.check('ACC-06','Cambio antes de verificar rechazado',self.send('/api/recuperacion/cambiar/',payload,client=c).status_code,403)
  self.check('ACC-06','Código auténtico del buzón QA aceptado',self.send('/api/recuperacion/verificar/',{'codigo':code},client=c).status_code,200)
  self.check('ACC-06','Clave débil rechazada por servidor',self.send('/api/recuperacion/cambiar/',{'password':'12345','confirmacion':'12345'},client=c).status_code,400)
  self.check('ACC-06','Confirmación distinta rechazada',self.send('/api/recuperacion/cambiar/',{**payload,'confirmacion':'OtraFixture!6752'},client=c).status_code,400)
  self.check('ACC-06','Guardar nueva clave',self.send('/api/recuperacion/cambiar/',payload,client=c).status_code,200)
  self.check('ACC-05','Código usado no puede verificarse de nuevo',self.send('/api/recuperacion/verificar/',{'codigo':code},client=c).status_code,400)
  self.check('ACC-06','Clave anterior rechazada',self.send('/api/login/',{'username':self.user.username,'password':'FixtureInicial!347652'},client=Client()).status_code,401)
  self.check('ACC-06','Nueva clave aceptada',self.send('/api/login/',{'username':self.user.username,'password':payload['password']},client=Client()).status_code,200)
 def test_ADM06_aislamiento_contexto(self):
  c=Client();c.force_login(self.user);context=c.get('/api/contexto/').json()['data']
  self.check('ADM-06','Solo delegación autorizada en selector',[v['id']for v in context['delegaciones_autorizadas']],[self.d.pk])
  self.check('ADM-06','Cambio a ámbito ajeno denegado',self.send('/api/contexto/cambiar-delegacion/',{'delegacion_id':self.d2.pk},client=c).status_code,403)
  self.user.delegaciones.add(self.d2)
  self.check('ADM-06','Cambio autorizado persistido',self.send('/api/contexto/cambiar-delegacion/',{'delegacion_id':self.d2.pk},client=c).status_code,200)
  self.check('ADM-06','Contexto activo actualizado',c.get('/api/contexto/').json()['data']['delegacion_activa']['id'],self.d2.pk)
  self.user.refresh_from_db();self.check('ADM-06','Delegación principal sin modificación',self.user.delegacion_id,self.d.pk)
 def attention(self,**extra):return {'solicitante':'Persona QA sin identificación real','tipo_atencion':self.tipo.pk,'sub_atencion':self.sub.pk,'descripcion':'Ensayo de atención',**extra}
 def test_ACT05_cuarta_gestion_y_par_invalido(self):
  for i in range(3):self.check('ACT-05','Gestión '+str(i+1)+' aceptada',self.send('/api/atenciones/',self.attention(caso_social='CASO-FICHA')).status_code,201)
  self.check('ACT-05','Cuarta gestión rechazada',self.send('/api/atenciones/',self.attention(caso_social='CASO-FICHA')).status_code,409)
  self.check('ACT-05','Caso conserva tres registros',Atencion.objects.filter(caso_social='CASO-FICHA').count(),3)
  other=TipoAtencion.objects.create(codigo='T2',nombre='Otro tipo')
  self.check('ACT-05','Par tipo subatención incoherente rechazado',self.send('/api/atenciones/',self.attention(tipo_atencion=other.pk)).status_code,400)
 def test_ACT06_idempotencia_y_contexto_servidor(self):
  c=Client();c.force_login(self.user);payload=self.attention(clave_idempotencia='FICHA-DOBLE')
  a=self.send('/api/atenciones/',payload,client=c);b=self.send('/api/atenciones/',payload,client=c)
  self.check('ACT-06','Primer envío crea',a.status_code,201);self.check('ACT-06','Reintento devuelve mismo registro',b.status_code,200)
  self.check('ACT-06','Mismo folio sin duplicación',a.json()['data']['folio'],b.json()['data']['folio'])
  self.check('ACT-06','Un solo registro persistido',Atencion.objects.count(),1)
  self.check('ACT-06','Responsable servidor',a.json()['data']['usuario'],self.user.pk)
  self.check('ACT-06','Delegación servidor',a.json()['data']['delegacion'],self.d.pk)
  self.check('ACT-06','Reutilización con datos distintos rechazada',self.send('/api/atenciones/',{**payload,'descripcion':'Cambio'},client=c).status_code,409)

class Result(unittest.TextTestResult):
 def stopTestRun(self):
  OUT.mkdir(exist_ok=True,parents=True)
  (OUT/'servidor-resultados.json').write_text(json.dumps({'entorno':'MySQL temporal test_sgr_eva2_local','correo':'Django locmem; sin SMTP externo; código obtenido del mensaje, no sustituido','fecha':datetime.now(ZoneInfo('America/Santiago')).isoformat(),'variantes':PROOFS,'tests':self.testsRun,'fallos':len(self.failures),'errores':len(self.errors)},ensure_ascii=False,indent=2,default=str))
  super().stopTestRun()
class Runner(DiscoverRunner):
 def get_test_runner_kwargs(self):return {**super().get_test_runner_kwargs(),'resultclass':Result}
