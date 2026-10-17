"""Reclamos y felicitaciones con datos reales del contexto y exportación protegida."""
from io import BytesIO
import json
from django.core.serializers.json import DjangoJSONEncoder
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from django.views.decorators.http import require_http_methods
from . import local_api as api,domain_api as domain
from .models import Reclamo,Usuario
FIELDS=['clase','fecha','vecino','contacto','ubicacion','situacion','profesional','medio_escrito','informado_a','evidencia']
def scoped(request):
    qs=Reclamo.objects.filter(delegacion=api.active_delegation(request)).select_related('autor','profesional','evidencia').order_by('-id')
    return qs if api.allowed(request.user,'reclamos','ver_equipo') else qs.filter(autor=request.user)
def data(obj):
    r=api.serialize(obj,FIELDS);r.update(version=obj.version,folio=obj.folio,autor=obj.autor_id,autor_nombre=obj.autor.get_full_name() or obj.autor.username,profesional_nombre=(obj.profesional.get_full_name() or obj.profesional.username) if obj.profesional_id else '',creado=obj.creado.isoformat(),archivo_url=f'/api/operacion/evidencias/{obj.evidencia_id}/archivo/' if obj.evidencia_id else None)
    return json.loads(json.dumps(r,cls=DjangoJSONEncoder))
@require_http_methods(['GET','POST','PATCH'])
def records(request,pk=None,action=None):
    denied=domain.gate(request,'reclamos','exportar' if action=='export' else 'ver' if request.method=='GET' else 'crear' if request.method=='POST' else 'editar')
    if denied:return denied
    if request.method=='PATCH' and not pk:return api.error('Selecciona un registro.',405)
    if action and request.method!='GET':return api.error('Método no permitido.',405)
    try:
        qs=scoped(request)
        if request.method=='GET':
            if pk:
                obj=qs.filter(pk=pk).first();return api.ok(data(obj)) if obj else api.error('Registro no encontrado.',404)
            if request.GET.get('search'):
                q=request.GET['search'];qs=qs.filter(Q(vecino__icontains=q)|Q(situacion__icontains=q)|Q(folio__icontains=q)|Q(ubicacion__icontains=q))
            if request.GET.get('clase'):qs=qs.filter(clase=request.GET['clase'])
            for name,op in [('desde','gte'),('hasta','lte')]:
                if request.GET.get(name):qs=qs.filter(**{'fecha__'+op:request.GET[name]})
            if action=='export':
                from openpyxl import Workbook
                from openpyxl.styles import Font,PatternFill
                book=Workbook();sheet=book.active;sheet.title='Reclamos y felicitaciones';sheet.append(['Folio','Fecha','Clase','Usuario / vecino','Contacto','Ubicación','Situación','Profesional involucrado','Medio escrito','Informado a','Código evidencia'])
                for obj in qs:
                    r=data(obj);sheet.append([r[k] for k in ['folio','fecha','clase','vecino','contacto','ubicacion','situacion','profesional_nombre','medio_escrito','informado_a','evidencia']])
                    for cell in sheet[sheet.max_row]:
                        if isinstance(cell.value,str):cell.data_type='s'
                for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='AD0000')
                sheet.freeze_panes='D2';sheet.auto_filter.ref=sheet.dimensions
                for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=25
                out=BytesIO();book.save(out);response=HttpResponse(out.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet');response['Content-Disposition']='attachment; filename="reclamos-felicitaciones.xlsx"';api.record_event(request,'reclamos',None,'exportar',None,{'cantidad':qs.count()});return response
            page=int(request.GET.get('page',1));pager=Paginator(qs,20)
            if page<1 or page>pager.num_pages:return api.error('Página inexistente.',404)
            return api.ok(results=[data(x) for x in pager.page(page)],count=pager.count,page=page,total_pages=pager.num_pages)
        incoming=api.body(request)
        if set(incoming)-set(FIELDS)-{'_version'}:return api.error('Campos no permitidos.')
        with transaction.atomic():
            obj=qs.select_for_update().filter(pk=pk).first() if pk else Reclamo(autor=request.user,delegacion=api.active_delegation(request))
            if obj is None:return api.error('Registro no encontrado.',404)
            if pk and request.method=='POST':return api.error('Método no permitido.',405)
            before=data(obj) if pk else None
            if pk:
                if obj.autor_id!=request.user.pk and not api.allowed(request.user,'reclamos','editar_equipo'):return api.error('No puedes editar un registro ajeno.',403)
                conflict=domain.validate_version(obj,incoming)
                if conflict:return conflict
            for field in FIELDS:
                if field not in incoming:continue
                value=incoming[field]
                if field in ['profesional','evidencia']:
                    if value in ['',None]:setattr(obj,field,None);continue
                    if isinstance(value,bool) or not str(value).isdigit():raise ValidationError('Referencia inválida.')
                    related=Usuario.objects.filter(pk=value,is_active=True).filter(Q(delegacion=obj.delegacion)|Q(delegaciones=obj.delegacion)).distinct().first() if field=='profesional' else domain.scoped(request,'evidencias').filter(pk=value).first() if api.allowed(request.user,'evidencias','ver') else None
                    if related is None:raise ValidationError('Referencia no autorizada.')
                    setattr(obj,field,related)
                elif field=='medio_escrito':
                    if not isinstance(value,bool):raise ValidationError('Medio escrito debe ser Sí o No.')
                    obj.medio_escrito=value
                else:
                    if not isinstance(value,str):raise ValidationError('Los campos deben ser texto.')
                    setattr(obj,field,value.strip())
            obj.full_clean()
            if obj.fecha>timezone.localdate():raise ValidationError('La fecha no puede ser futura.')
            if pk:obj.version+=1
            obj.save();api.record_event(request,'reclamos',obj.pk,'editar' if pk else 'crear',before,data(obj))
        response=api.ok(data(obj));response.status_code=200 if pk else 201;return response
    except (ValidationError,ValueError,TypeError) as ex:return api.error('; '.join(ex.messages) if isinstance(ex,ValidationError) else 'Datos inválidos.')

@require_http_methods(['GET'])
def options(request):
    denied=domain.gate(request,'reclamos','ver')
    if denied:return denied
    d=api.active_delegation(request)
    users=Usuario.objects.filter(is_active=True).filter(Q(delegacion=d)|Q(delegaciones=d)).distinct()
    return api.ok(responsables=[{'id':u.pk,'nombre':u.get_full_name() or u.username} for u in users])
