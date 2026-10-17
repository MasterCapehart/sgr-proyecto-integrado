"""API del incremento local; correo externo solo con SMTP configurado."""
import json
import logging
import hashlib
import secrets
import uuid
from datetime import timedelta
from io import BytesIO
from decimal import Decimal
from smtplib import SMTPException
from django.conf import settings
from django.contrib.auth import authenticate, login, logout, update_session_auth_hash
from django.contrib.auth.hashers import make_password, check_password
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError, ImproperlyConfigured
from django.core.mail import send_mail
from django.core.paginator import Paginator
from django.db import transaction, IntegrityError
from django.db.models import Q
from django.http import JsonResponse, HttpResponse
from django.core.serializers.json import DjangoJSONEncoder
from django.utils import timezone
from django.views.decorators.csrf import ensure_csrf_cookie
from django.views.decorators.http import require_http_methods
from .models import (Delegacion, Rol, Usuario, TipoAtencion, SubAtencion, ItemMeta,
 Cargo, ConfiguracionMeta, Atencion, AuditoriaDerivacion, RecuperacionAcceso, CasoSocial, EventoLocal, AccesoFallido)

logger = logging.getLogger(__name__)

RESOURCES = {
 'cargos': (Cargo, ['codigo','nombre','area','planilla','estado']),
 'delegaciones': (Delegacion, ['codigo','nombre','direccion','comuna','estado']),
 'roles': (Rol, ['nombre','descripcion','alcance','permisos','estado']),
 'usuarios': (Usuario, ['username','first_name','last_name','email','rut','rol','cargo','delegacion','delegaciones','is_active']),
 'tipos-atencion': (TipoAtencion, ['codigo','nombre','descripcion','estado']),
 'subatenciones': (SubAtencion, ['codigo','nombre','descripcion','tipo_atencion','estado']),
 'items-metas': (ItemMeta, ['codigo','nombre','descripcion','unidad_medida','cargos','estado']),
 'configuraciones-metas': (ConfiguracionMeta, ['item','periodo','meta_esperada','ponderacion','estado']),
 'atenciones': (Atencion, ['folio','solicitante','tipo_atencion','sub_atencion','descripcion','delegacion','delegacion_responsable','usuario','fecha_registro','caso_social','numero_gestion','rut_solicitante','telefono_solicitante','requiere_visita','fecha_visita','fecha_entrega_informe','fecha_entrega_beneficio','gestion_realizada','fecha_visita_programada','fecha_visita_realizada','tipo_gestion','compromiso','estado']),
}
ACTIONS = ['ver','crear','editar','desactivar','exportar']
from .domain_permissions import DOMAIN_PERMISSIONS
PERMISSIONS = [f'{resource}.{action}' for resource in RESOURCES for action in ACTIONS]+['auditoria.ver']+DOMAIN_PERMISSIONS

def ok(data=None, **extra):
    payload={'success':True, **extra}
    if data is not None: payload['data']=data
    return JsonResponse(payload)

def error(message, status=400, details=None):
    return JsonResponse({'success':False,'error':{'code':str(status),'message':message,'details':details or {}}},status=status)

def body(request):
    try: data=json.loads(request.body or '{}') if request.content_type=='application/json' else request.POST.dict()
    except (ValueError, UnicodeDecodeError): raise ValidationError('JSON inválido.')
    if not isinstance(data,dict): raise ValidationError('Se requiere un objeto JSON.')
    return data

def allowed(user, resource, action='ver'):
    return user.is_authenticated and user.is_active and (user.is_superuser or
        bool(user.rol_id and user.rol.estado and f'{resource}.{action}' in user.rol.permisos))

def delegations(user):
    qs=Delegacion.objects.filter(estado=True)
    if user.is_superuser: return qs
    return qs.filter(Q(usuarios_autorizados=user)|Q(pk=user.delegacion_id)).distinct()

def active_delegation(request):
    qs=delegations(request.user)
    return qs.filter(pk=request.session.get('delegacion_activa') or request.user.delegacion_id).first() or qs.first()

def serialize(obj, fields):
    result={'id':obj.pk}
    if hasattr(obj,'version'): result['version']=obj.version
    for name in fields:
        field=obj._meta.get_field(name)
        if field.many_to_many: value=list(getattr(obj,name).values_list('pk',flat=True))
        elif field.is_relation: value=getattr(obj,name+'_id')
        else: value=getattr(obj,name)
        result[name]=value
    if isinstance(obj,SubAtencion): result['tipo_padre_display']=f'{obj.tipo_atencion.codigo} · {obj.tipo_atencion.nombre}'
    if isinstance(obj,Atencion):
        snapshot=obj.catalogos_snapshot or {}
        result.update(usuario_nombre=obj.usuario.get_full_name() or obj.usuario.username,tipo_nombre=snapshot.get('tipo_atencion',{}).get('nombre',obj.tipo_atencion.nombre),subatencion_nombre=snapshot.get('sub_atencion',{}).get('nombre',obj.sub_atencion.nombre),catalogos_snapshot=snapshot or None,historia_catalogos='Captura posterior; etiqueta original no garantizada' if snapshot.get('observada') else 'Captura al registrar' if snapshot else 'Sin captura histórica; etiqueta actual')
    if isinstance(obj,Usuario): result['estado']=obj.is_active; result['nombre']=obj.get_full_name() or obj.username
    return result

@ensure_csrf_cookie
@require_http_methods(['GET'])
def csrf(request):
    from django.middleware.csrf import get_token
    return ok({"csrfToken": get_token(request)})

@require_http_methods(['POST'])
def access(request):
    try: data=body(request)
    except ValidationError as ex: return error(ex.messages[0])
    if not isinstance(data.get('username',''),str) or not isinstance(data.get('password',''),str):return error('Usuario y contraseña deben ser texto.')
    identifier=data.get('username','').strip()
    if '@' in identifier:
        found=Usuario.objects.filter(email__iexact=identifier).first()
        identifier=found.username if found else identifier
    # Límite persistente por identificador e IP, además del límite de sesión.
    account_hash=hashlib.sha256(identifier.casefold().encode()).hexdigest()
    address_hash=hashlib.sha256(request.META.get('REMOTE_ADDR','local').encode()).hexdigest()
    if AccesoFallido.objects.filter(creado__gt=timezone.now()-timedelta(minutes=5)).filter(Q(cuenta_hash=account_hash)|Q(direccion_hash=address_hash)).count()>=5: return error('Espera cinco minutos antes de volver a intentar.',429)
    # Sin revelar si la cuenta existe.
    now=timezone.now().timestamp(); attempts=request.session.get('login_attempts',[])
    attempts=[t for t in attempts if now-t<300]
    if len(attempts)>=5: return error('Espera cinco minutos antes de volver a intentar.',429)
    user=authenticate(request,username=identifier,password=data.get('password',''))
    if not user or (not user.is_superuser and (not user.rol_id or not user.rol.estado)):
        AccesoFallido.objects.create(cuenta_hash=account_hash,direccion_hash=address_hash)
        request.session['login_attempts']=attempts+[now]
        return error('Credenciales incorrectas o cuenta sin acceso.',401)
    login(request,user); request.session['login_attempts']=[]
    record_event(request,'acceso',user.pk,'iniciar_sesion',None,{'sesion':'iniciada'})
    return ok({'usuario':user.get_full_name() or user.username})

@require_http_methods(['POST'])
def exit_access(request):
    if request.user.is_authenticated:record_event(request,'acceso',request.user.pk,'cerrar_sesion',{'sesion':'activa'},{'sesion':'cerrada'})
    logout(request); return ok()

@require_http_methods(['GET'])
def context(request):
    if not request.user.is_authenticated: return error('Inicia sesión.',401)
    u=request.user
    if not u.is_superuser and (not u.rol_id or not u.rol.estado): return error('Tu rol no está habilitado.',403)
    d=active_delegation(request)
    ds=list(delegations(u).values('id','codigo','nombre'))
    user={'id':u.pk,'nombre':u.get_full_name() or u.username,'username':u.username,'nombre_completo':u.get_full_name() or u.username,'display':u.get_full_name() or u.username,'email':u.email,'rut':u.rut or '', 'cargo':u.cargo_id,'cargo_nombre':u.cargo.nombre if u.cargo_id else None,'area':u.cargo.area if u.cargo_id else None,'rol':u.rol.nombre if u.rol_id else 'Administrador'}
    delegation={'id':d.pk,'codigo':d.codigo,'nombre':d.nombre} if d else None
    return ok({'usuario':user,'rol':user['rol'],'delegacion_activa':delegation,'delegaciones_autorizadas':ds,
               'delegaciones':ds,'delegacion':delegation,'fecha_servidor':timezone.now().isoformat(),'fecha_display':timezone.localtime().strftime('%d/%m/%Y %H:%M'),'permisos':PERMISSIONS if u.is_superuser else (u.rol.permisos if u.rol_id else [])})

@require_http_methods(['POST'])
def change_context(request):
    if not request.user.is_authenticated: return error('Inicia sesión.',401)
    try: data=body(request); pk=data.get('delegacion_id',data.get('delegacion'))
    except ValidationError as ex: return error(ex.messages[0])
    if not str(pk).isdigit() or not delegations(request.user).filter(pk=pk).exists(): return error('Delegación no autorizada.',403)
    previous=active_delegation(request)
    request.session['delegacion_activa']=int(pk)
    record_event(request,'contexto',request.user.pk,'cambiar_delegacion',{'delegacion':previous.pk if previous else None},{'delegacion':int(pk)})
    response=context.__wrapped__(request)
    payload=json.loads(response.content); payload['data'].update(payload['data']['delegacion_activa']);return JsonResponse(payload)

@require_http_methods(['GET'])
def recovery_status(request):
    identifier=request.session.get('recovery_id')
    rec=RecuperacionAcceso.objects.filter(identificador=identifier,verificado=True,consumido=False,vence__gt=timezone.now(),intentos__lt=5).first() if identifier else None
    return ok({'verificado':True}) if rec else error('Primero verifica un código vigente.',403)

@require_http_methods(['POST'])
def recovery(request, step='solicitar'):
    try:
        data=body(request)
        if step=='solicitar':
            email=str(data.get('email','')).strip().lower()
            from django.core.validators import validate_email
            validate_email(email)
            now=timezone.now(); last=request.session.get('recovery_sent',0)
            if now.timestamp()-last<60: return error('Espera 60 segundos para reenviar.',429)
            identifier=uuid.uuid4(); user=Usuario.objects.filter(email__iexact=email,is_active=True).first()
            if user:
                with transaction.atomic():
                    Usuario.objects.select_for_update().get(pk=user.pk)
                    recent=RecuperacionAcceso.objects.filter(usuario=user,creado__gt=now-timedelta(seconds=60)).exists()
                    if not recent:
                        code=f'{secrets.randbelow(1000000):06d}'
                        recovery_record=RecuperacionAcceso.objects.create(usuario=user,identificador=identifier,codigo_hash=make_password(code),vence=now+timedelta(minutes=10))
                        try:
                            sent=send_mail(
                                'SGR: código de recuperación de contraseña',
                                f'Tu código de verificación es: {code}\n\n'
                                'Vence en 10 minutos y solo puede utilizarse una vez.\n'
                                'No compartas este código. Si no solicitaste el cambio, ignora este mensaje.\n\n'
                                'Sistema de Gestión de Resultados (SGR)',
                                settings.DEFAULT_FROM_EMAIL,[email],fail_silently=False,
                            )
                            if sent != 1:
                                raise SMTPException('El servicio no aceptó el mensaje.')
                        except (SMTPException, OSError, ValueError, ImproperlyConfigured) as exc:
                            recovery_record.consumido=True
                            recovery_record.save(update_fields=['consumido'])
                            # No registrar destinatarios, códigos, credenciales ni respuestas SMTP.
                            logger.error('No se pudo entregar un código SGR. Tipo de fallo: %s',type(exc).__name__)
                        else:
                            RecuperacionAcceso.objects.filter(usuario=user,consumido=False).exclude(pk=recovery_record.pk).update(consumido=True)
            request.session['recovery_id']=str(identifier); request.session['recovery_sent']=now.timestamp()
            return ok(message='Solicitud recibida. Si la cuenta está habilitada, recibirás un código. Revisa también la carpeta de correo no deseado. Si no llega, espera 60 segundos y vuelve a solicitarlo.')
        identifier=request.session.get('recovery_id')
        with transaction.atomic():
            rec=RecuperacionAcceso.objects.select_for_update().filter(identificador=identifier).first() if identifier else None
            if not rec or rec.consumido or rec.vence<=timezone.now() or rec.intentos>=5: return error('Código inválido, agotado o vencido.')
            if step=='verificar':
                rec.intentos+=1
                valid=check_password(str(data.get('codigo',data.get('code',''))),rec.codigo_hash)
                rec.verificado=valid; rec.save()
                return ok() if valid else error('Código inválido, agotado o vencido.')
            if not rec.verificado: return error('Primero verifica el código.',403)
            password=data.get('password','')
            if not isinstance(password,str) or not isinstance(data.get('confirmacion',password),str):return error('La contraseña y su confirmación deben ser texto.')
            validate_password(password,rec.usuario)
            if password!=data.get('confirmacion',password): return error('Las contraseñas no coinciden.')
            user=Usuario.objects.select_for_update().get(pk=rec.usuario_id)
            previous_version=user.version
            user.set_password(password); user.codigo_recuperacion=None; user.version+=1; user.save()
            EventoLocal.objects.create(usuario=user,recurso='usuarios',registro=user.pk,accion='cambiar_clave',antes={'credencial':'anterior','version':previous_version},despues={'credencial':'renovada','version':user.version},origen=request.META.get('REMOTE_ADDR'),delegacion=user.delegacion)
            RecuperacionAcceso.objects.filter(usuario=user,consumido=False).update(consumido=True)
            request.session.pop('recovery_id',None)
        return ok(message='Contraseña actualizada. Inicia sesión con tu nueva contraseña.')
    except ValidationError as ex: return error('Revisa los datos.',details={'campos':ex.messages})

@require_http_methods(['GET'])
def options(request):
    if not request.user.is_authenticated: return error('Inicia sesión.',401)
    if not any(allowed(request.user,r) for r in RESOURCES): return error('Sin permisos.',403)
    from .matrix_profiles import PROFILES
    from .excel_support import GESTIONES
    result={'gestiones':GESTIONES,'planillas':[{'id':k,'nombre':v['nombre']} for k,v in PROFILES.items() if k!='GENERAL'],'cargos':list(Cargo.objects.filter(estado=True).values('id','codigo','nombre','area')),'delegaciones':list(delegations(request.user).values('id','nombre','codigo')),
            'tipos-atencion':list(TipoAtencion.objects.filter(estado=True).values('id','nombre','codigo')),
            'items-metas':list(ItemMeta.objects.filter(estado=True).values('id','nombre','codigo'))}
    if allowed(request.user,'usuarios','crear') or allowed(request.user,'usuarios','editar'):
        result['roles']=list(Rol.objects.filter(estado=True).values('id','nombre'))
    if allowed(request.user,'roles'): result['permisos']=PERMISSIONS
    return ok(result)

def validate_data(obj, data, resource, request, creating=False):
    fields=RESOURCES[resource][1]; incoming=dict(data)
    if resource=='usuarios' and 'estado' in incoming: incoming['is_active']=incoming.pop('estado')
    valid_fields=set(fields)|({'password'} if resource=='usuarios' else set())
    if set(incoming)-valid_fields: raise ValidationError('Hay campos desconocidos o no editables.')
    for key,value in incoming.items():
        if key in ['password','delegaciones','cargos']: continue
        f=obj._meta.get_field(key)
        if f.is_relation:
            if value in [None,'']: setattr(obj,key,None)
            else:
                if not str(value).isdigit(): raise ValidationError({key:'Selecciona una opción válida.'})
                related=f.related_model.objects.select_for_update().filter(pk=value).first()
                if not related or not getattr(related,'estado',True): raise ValidationError({key:'Registro relacionado inexistente o inactivo.'})
                setattr(obj,key,related)
        else:
            if f.get_internal_type() in ['CharField','TextField','EmailField'] and value is not None and not isinstance(value,str): raise ValidationError({key:'Se requiere texto.'})
            if f.get_internal_type()=='BooleanField' and not isinstance(value,bool): raise ValidationError({key:'Debe ser verdadero o falso.'})
            setattr(obj,key,value.strip() if isinstance(value,str) else value)
    missing={key:['Campo obligatorio.'] for key in ['codigo','nombre'] if hasattr(obj,key) and not getattr(obj,key)}
    if missing: raise ValidationError(missing)
    for key in ['codigo','nombre']:
        if hasattr(obj,key):
            value=getattr(obj,key)
            if not value: raise ValidationError({key:'Campo obligatorio.'})
            if type(obj).objects.filter(**{key+'__iexact':value}).exclude(pk=obj.pk).exists(): raise ValidationError({key:'Ya existe un registro con ese valor.'})
    if isinstance(obj,Rol):
        if not isinstance(obj.permisos,list) or any(p not in PERMISSIONS for p in obj.permisos): raise ValidationError({'permisos':'Permisos inválidos.'})
        if not request.user.is_superuser and (obj.alcance=='global' or set(obj.permisos)-set(request.user.rol.permisos)): raise ValidationError('No puedes otorgar permisos superiores a los propios.')
        if obj.pk and obj.pk==request.user.rol_id: raise ValidationError('No puedes modificar tu propio rol en esta operación.')
        if obj.pk and not obj.estado and obj.usuario_set.filter(is_active=True).exists(): raise ValidationError('El rol tiene usuarios activos.')
    if isinstance(obj,Usuario):
        from django.core.validators import validate_email
        obj.rut=obj.rut or None
        validate_email(obj.email)
        if not obj.first_name.strip(): raise ValidationError({'first_name':'Nombre obligatorio.'})
        if Usuario.objects.filter(email__iexact=obj.email).exclude(pk=obj.pk).exists(): raise ValidationError({'email':'El correo ya está registrado.'})
        if not obj.rol_id: raise ValidationError({'rol':'Selecciona un rol.'})
        if obj.is_active and not obj.rol.estado:raise ValidationError({'rol':'No se puede activar una cuenta con un rol inactivo.'})
        ds=incoming.get('delegaciones', list(obj.delegaciones.values_list('pk',flat=True)) if obj.pk else [])
        if not isinstance(ds,list) or any(not str(d).isdigit() for d in ds): raise ValidationError({'delegaciones':'Selecciona delegaciones válidas.'})
        ds=list(set(int(d) for d in ds))
        if not ds or not obj.delegacion_id or obj.delegacion_id not in ds: raise ValidationError({'delegaciones':'Incluye la delegación principal entre las autorizadas.'})
        if len(ds)!=delegations(request.user).filter(pk__in=ds).count(): raise ValidationError('No puedes asignar delegaciones no autorizadas.')
        if not request.user.is_superuser and (obj.is_superuser or obj.is_staff or obj.rol.alcance=='global' or set(obj.rol.permisos)-set(request.user.rol.permisos)): raise ValidationError('No puedes administrar una cuenta con privilegios superiores.')
        if obj.pk==request.user.pk and not obj.is_active: raise ValidationError('No puedes desactivar tu propia cuenta.')
        password=incoming.get('password')
        if password is not None and not isinstance(password,str):raise ValidationError({'password':'La contraseña debe ser texto.'})
        if creating and not password: raise ValidationError({'password':'Contraseña obligatoria.'})
        if password: validate_password(password,obj); obj.set_password(password)
        obj._delegaciones=ds
    if isinstance(obj,ItemMeta) and 'cargos' in incoming:
        ids=incoming['cargos']
        if not isinstance(ids,list) or any(isinstance(i,bool) or not str(i).isdigit() for i in ids):raise ValidationError({'cargos':'Selecciona cargos válidos.'})
        ids=list(set(int(i) for i in ids))
        if Cargo.objects.filter(pk__in=ids,estado=True).count()!=len(ids):raise ValidationError({'cargos':'Hay cargos inexistentes o inactivos.'})
        obj._cargos=ids
    if isinstance(obj,Usuario) and obj.is_active and obj.cargo_id and not obj.cargo.estado:raise ValidationError({'cargo':'El cargo está inactivo.'})
    if isinstance(obj,SubAtencion):
        if not obj.tipo_atencion_id:raise ValidationError({'tipo_atencion':'Selecciona el tipo padre.'})
        if obj.estado and not obj.tipo_atencion.estado:raise ValidationError({'tipo_atencion':'No se puede activar una subatención cuyo tipo está inactivo.'})
    if isinstance(obj,ConfiguracionMeta):
        if obj.estado and obj.item_id and not obj.item.estado:raise ValidationError({'item':'No se puede activar una meta de un ítem inactivo.'})
        if obj.item_id:
            # Una misma fila de ítem serializa creaciones y cambios concurrentes.
            # No altera ni elimina configuraciones históricas preexistentes.
            ItemMeta.objects.select_for_update().get(pk=obj.item_id)
            if ConfiguracionMeta.objects.filter(item_id=obj.item_id,periodo__iexact=obj.periodo).exclude(pk=obj.pk).exists():
                raise ValidationError({'periodo':'Ya existe una meta para este ítem y período. Edita la configuración existente.'})
        if Decimal(str(obj.meta_esperada))<=0 or not 0<Decimal(str(obj.ponderacion))<=100: raise ValidationError('Meta positiva y ponderación entre 0 y 100.')
    if obj.pk and hasattr(obj,'estado') and not obj.estado:
        relations=[]
        if isinstance(obj,Delegacion): relations=[obj.usuario_set.filter(is_active=True), obj.usuarios_autorizados.filter(is_active=True),obj.atenciones_origen.all(),obj.atenciones_responsable.all()]
        if isinstance(obj,TipoAtencion): relations=[obj.subatenciones.filter(estado=True),obj.atenciones.all()]
        if isinstance(obj,SubAtencion): relations=[obj.atenciones.all()]
        if isinstance(obj,Cargo): relations=[obj.usuario_set.filter(is_active=True),obj.items.filter(estado=True)]
        if isinstance(obj,ItemMeta): relations=[obj.configuraciones_meta.filter(estado=True)]
        if any(q.exists() for q in relations): raise ValidationError('No se puede desactivar: existen registros relacionados.')
    obj.full_clean(exclude=['password'] if isinstance(obj,Usuario) else None)

@require_http_methods(['GET','POST','PUT','PATCH','DELETE'])
def crud(request, resource, pk=None, action=None):
    if resource not in RESOURCES: return error('Recurso inexistente.',404)
    if not request.user.is_authenticated: return error('Inicia sesión.',401)
    verb='exportar' if action=='export' else ('ver' if request.method=='GET' else 'desactivar' if request.method=='DELETE' or action=='toggle-estado' else 'crear' if request.method=='POST' else 'editar')
    if (action=='toggle-estado' and request.method!='POST') or (pk and request.method=='POST' and not action): return error('Método no permitido.',405)
    if not allowed(request.user,resource,verb):
        record_event(request,resource,pk,'denegar_'+verb,None,None)
        return error('No tienes permiso para esta operación.',403)
    model,fields=RESOURCES[resource]; qs=model.objects.all()
    if resource=='delegaciones' and not request.user.is_superuser: qs=qs.filter(pk__in=delegations(request.user))
    if resource=='usuarios' and not request.user.is_superuser: qs=qs.filter(is_superuser=False,is_staff=False,delegacion__in=delegations(request.user))
    if resource=='atenciones': qs=qs.filter(Q(delegacion=active_delegation(request))|Q(delegacion_responsable=active_delegation(request)))
    try:
        if request.method=='GET':
            if pk:
                obj=qs.filter(pk=pk).first()
                return ok(serialize(obj,fields)) if obj else error('Registro no encontrado.',404)
            search=request.GET.get('search','').strip(); state=request.GET.get('estado','activos')
            if state not in ['activos','inactivos','todos']: return error('Filtro de estado inválido.')
            statefield='is_active' if resource=='usuarios' else 'estado'
            if state!='todos': qs=qs.filter(**{statefield:state=='activos'})
            if search:
                query=Q(tipo_atencion__nombre__icontains=search) if resource=='subatenciones' else Q()
                for name in fields:
                    f=model._meta.get_field(name)
                    if f.get_internal_type() in ['CharField','TextField','EmailField']: query|=Q(**{name+'__icontains':search})
                qs=qs.filter(query)
            for parent in ['tipo_atencion','item','rol','delegacion']:
                if parent in fields and request.GET.get(parent):
                    if not request.GET[parent].isdigit(): return error('Filtro relacionado inválido.')
                    qs=qs.filter(**{parent+'_id':request.GET[parent]})
            ordering=request.GET.get('ordering','-id')
            safe=['id']+[f for f in fields if not model._meta.get_field(f).many_to_many and not model._meta.get_field(f).is_relation and f!='permisos']
            if ordering.lstrip('-') not in safe: return error('Ordenamiento inválido.')
            qs=qs.order_by(ordering,'id').distinct()
            if action=='export':
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill
                book=Workbook(); sheet=book.active; sheet.title='Resultados SGR'; sheet.append(['ID']+[('Nombre del Ítem' if resource=='items-metas' and f=='nombre' else 'Tipo de Atención Padre' if resource=='subatenciones' and f=='tipo_atencion' else f) for f in fields])
                for obj in qs.iterator():
                    record=serialize(obj,fields); row=[record.get(f) for f in ['id']+fields]
                    sheet.append([str(v) if isinstance(v,(list,dict)) or hasattr(v,'tzinfo') else v for v in row])
                    for cell in sheet[sheet.max_row]:
                        if isinstance(cell.value,str): cell.data_type='s'
                for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='9F1324')
                sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
                for col in sheet.columns: sheet.column_dimensions[col[0].column_letter].width=min(50,max(14,max(len(str(c.value or '')) for c in col)+2))
                stream=BytesIO();book.save(stream)
                record_event(request,resource,None,'exportar',None,None)
                return HttpResponse(stream.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{resource}.xlsx"'})
            try: size=int(request.GET.get('page_size',10)); page=int(request.GET.get('page',1))
            except ValueError: return error('Paginación inválida.')
            if not 1<=size<=100 or page<1: return error('Paginación fuera de rango.')
            paginator=Paginator(qs,size)
            if page>paginator.num_pages: return error('Página inexistente.',404)
            part=paginator.page(page)
            from .kpis import resource_kpis
            extra={'kpis':resource_kpis(qs,resource)} if resource=='atenciones' else {}
            return ok(results=[serialize(o,fields) for o in part], count=paginator.count,page=page,page_size=size,total_pages=paginator.num_pages,**extra)
        if resource=='atenciones':
            if request.method!='POST' or pk or action: return error('El registro de atención no permite modificaciones retrospectivas.',405)
            return create_attention(request,body(request))
        with transaction.atomic():
            obj=qs.select_for_update().filter(pk=pk).first() if pk else model()
            if obj is None: return error('Registro no encontrado.',404)
            if pk and not request.user.is_superuser:
                target_role=obj if isinstance(obj,Rol) else obj.rol if isinstance(obj,Usuario) else None
                if target_role and (target_role.alcance=='global' or set(target_role.permisos)-set(request.user.rol.permisos)): return error('No puedes administrar privilegios superiores.',403)
            if pk is None and request.method!='POST': return error('Falta el registro.',405)
            data=body(request)
            expected=data.pop('_version',None)
            if pk and expected is None: return error('Recarga el registro antes de guardar: falta su versión.',428)
            if pk and (not isinstance(expected,int) or isinstance(expected,bool) or expected!=obj.version): return error('Otra persona modificó este registro. Recarga y revisa los cambios antes de guardar.',409)
            previous=snapshot(obj,fields) if pk else None
            if request.method=='DELETE' or action=='toggle-estado':
                if not pk: return error('Falta el registro.')
                statefield='is_active' if resource=='usuarios' else 'estado'
                data={statefield:False if request.method=='DELETE' else not getattr(obj,statefield)}
            if pk and ('estado' in data or 'is_active' in data) and not allowed(request.user,resource,'desactivar'): return error('Sin permiso para cambiar estado.',403)
            from . import catalog_history
            if pk:
                catalog_history.record(request,resource,obj,True)
                catalog_history.freeze_references(resource,obj)
            validate_data(obj,data,resource,request,creating=not pk)
            if pk: obj.version+=1
            obj.save()
            if hasattr(obj,'_delegaciones'): obj.delegaciones.set(obj._delegaciones)
            if hasattr(obj,'_cargos'): obj.cargos.set(obj._cargos)
            catalog_history.record(request,resource,obj)
            record_event(request,resource,obj.pk,verb,previous,snapshot(obj,fields))
            response=ok(serialize(obj,fields),message='Cambio guardado.')
            if not pk: response.status_code=201
            return response
    except ValidationError as ex:
        if resource=='atenciones' and request.method=='POST' and request.content_type=='application/json':
            try:replay=body(request)
            except ValidationError:return error('JSON inválido.')
            if replay.get('clave_idempotencia') and Atencion.objects.filter(clave_idempotencia=replay['clave_idempotencia']).exists(): return create_attention(request,replay)
        return error('Revisa los campos.',details=getattr(ex,'message_dict',{'campos':ex.messages}))
    except (ValueError,TypeError): return error('Formato de datos inválido.')
    except IntegrityError:
        if resource=='atenciones' and request.method=='POST' and request.content_type=='application/json':
            try:replay=body(request)
            except ValidationError:return error('JSON inválido.')
            if replay.get('clave_idempotencia') and Atencion.objects.filter(clave_idempotencia=replay['clave_idempotencia']).exists(): return create_attention(request,replay)
        return error('Conflicto: registro duplicado o relacionado.',409)

def create_attention(request,data):
    from .excel_support import SOCIAL_FIELDS, validate_social
    allowed_keys=set(SOCIAL_FIELDS)|{'solicitante','tipo_atencion','sub_atencion','descripcion','delegacion_responsable','caso_social','clave_idempotencia','motivo_derivacion','compromiso'}
    if set(data)-allowed_keys: return error('Hay campos no editables.')
    d=active_delegation(request)
    if not d: return error('No tienes delegación autorizada.',403)
    key=data.get('clave_idempotencia') or uuid.uuid4().hex
    if not isinstance(key,str) or len(key)>64: return error('Clave de idempotencia inválida.')
    existing=Atencion.objects.filter(clave_idempotencia=key).first()
    if existing:
        if existing.usuario_id!=request.user.pk or existing.delegacion_id!=d.pk: return error('Clave ya utilizada.',409)
        for field in ['solicitante','descripcion','caso_social','tipo_atencion_id','sub_atencion_id','delegacion_responsable_id','compromiso_id']:
            source=field.removesuffix('_id')
            if source in data and field.endswith('_id') and str(getattr(existing,field))!=str(data[source]): return error('La clave corresponde a otra solicitud.',409)
            if field in data and (getattr(existing,field) or '')!=((str(data[field]).strip().upper() if field=='caso_social' else str(data[field]).strip()) if data[field] else ''): return error('La clave corresponde a otra solicitud.',409)
        from copy import copy
        normalized=copy(existing)
        try:validate_social(normalized,data)
        except ValidationError as ex:return error('; '.join(ex.messages))
        for field in SOCIAL_FIELDS:
            if field in data and str(getattr(existing,field) or '')!=str(getattr(normalized,field) or ''):return error('La clave corresponde a otra solicitud.',409)
        return ok(serialize(existing,RESOURCES['atenciones'][1]),message='Registro ya guardado.',is_duplicate_prevented=True)
    obj=Atencion(folio='AT-'+uuid.uuid4().hex.upper(),usuario=request.user,delegacion=d,clave_idempotencia=key)
    if any(not isinstance(data.get(field,''),str) for field in ['solicitante','descripcion','caso_social'] if data.get(field) is not None): return error('Solicitante, descripción y caso deben ser texto.')
    for field in ['solicitante','descripcion','caso_social']: setattr(obj,field,str(data.get(field) or '').strip() or None)
    if not obj.solicitante or not obj.descripcion: return error('Solicitante y descripción son obligatorios.',details={f:['Campo obligatorio.'] for f in ['solicitante','descripcion'] if not getattr(obj,f)})
    for field,model in [('tipo_atencion',TipoAtencion),('sub_atencion',SubAtencion)]:
        value=data.get(field)
        if not str(value).isdigit(): return error('Selecciona tipo y subatención.')
        related=model.objects.filter(pk=value,estado=True).first()
        if not related: return error('Catálogo inexistente o inactivo.',details={field:['Catálogo inexistente o inactivo.']})
        setattr(obj,field,related)
    destination=data.get('delegacion_responsable') or d.pk
    obj.delegacion_responsable=delegations(request.user).filter(pk=destination).first()
    if not obj.delegacion_responsable: return error('Destino no autorizado.',403)
    if data.get('compromiso') not in ['',None]:
        if isinstance(data['compromiso'],bool) or not str(data['compromiso']).isdigit():return error('Compromiso inválido.')
        from . import domain_api
        if not allowed(request.user,'compromisos','ver'):return error('No puedes vincular compromisos.',403)
        obj.compromiso=domain_api.scoped(request,'compromisos').filter(pk=data['compromiso']).first()
        if obj.compromiso is None:return error('Compromiso no autorizado.')
    try:validate_social(obj,data)
    except ValidationError as ex:return error('; '.join(ex.messages))
    with transaction.atomic():
        if obj.caso_social:
            obj.caso_social=obj.caso_social.upper()
            case,_=CasoSocial.objects.get_or_create(codigo=obj.caso_social,defaults={'solicitante':' '.join(obj.solicitante.casefold().split())})
            case=CasoSocial.objects.select_for_update().get(pk=case.pk)
            prior=Atencion.objects.filter(caso_social=obj.caso_social)
            if prior.exists() and not prior.filter(Q(delegacion=d)|Q(delegacion_responsable=d)).exists():return error('El código de caso no está disponible en esta delegación.',409)
            previous_rut=prior.exclude(rut_solicitante='').values_list('rut_solicitante',flat=True).first()
            if previous_rut and obj.rut_solicitante and previous_rut!=obj.rut_solicitante:return error('El RUN no coincide con el del expediente.',409)
            if case.solicitante!=' '.join(obj.solicitante.casefold().split()): return error('El caso pertenece a otra persona.',409)
            obj.numero_gestion=Atencion.objects.filter(caso_social=obj.caso_social).count()+1
            if obj.numero_gestion>3: return error('El caso ya tiene tres gestiones.',409)
        from .excel_support import validate_stage
        validate_stage(obj)
        from . import catalog_history
        obj.tipo_atencion=TipoAtencion.objects.select_for_update().get(pk=obj.tipo_atencion_id)
        obj.sub_atencion=SubAtencion.objects.select_for_update().get(pk=obj.sub_atencion_id)
        if not obj.tipo_atencion.estado or not obj.sub_atencion.estado:raise ValidationError('El catálogo dejó de estar activo.')
        obj.catalogos_snapshot=catalog_history.attention(obj)
        obj.full_clean();obj.save()
        if obj.delegacion_responsable_id!=d.pk: AuditoriaDerivacion.objects.create(atencion=obj,delegacion_origen=d,delegacion_destino=obj.delegacion_responsable,usuario_responsable=request.user,motivo=data.get('motivo_derivacion','Derivación institucional autorizada'))
        record_event(request,'atenciones',obj.pk,'crear',None,snapshot(obj,RESOURCES['atenciones'][1]))
    response=ok(serialize(obj,RESOURCES['atenciones'][1]),message='Atención registrada.');response.status_code=201;return response


def snapshot(obj,fields):
    # Lista explícita de campos de la API: nunca password, tokens o códigos de recuperación.
    return json.loads(json.dumps(serialize(obj,fields),cls=DjangoJSONEncoder))

def record_event(request,resource,pk,action,before,after):
    EventoLocal.objects.create(usuario=request.user,recurso=resource,registro=pk,accion=action,
        antes=before,despues=after,origen=request.META.get('REMOTE_ADDR'),delegacion=active_delegation(request))

@require_http_methods(['GET'])
def audit(request):
    if not request.user.is_authenticated:return error('Inicia sesión.',401)
    if not allowed(request.user,'auditoria'):return error('Sin permiso de auditoría.',403)
    qs=EventoLocal.objects.select_related('usuario','delegacion').filter(delegacion=active_delegation(request)).order_by('-id')
    search=request.GET.get('search','').strip()
    if search:qs=qs.filter(Q(recurso__icontains=search)|Q(accion__icontains=search)|Q(usuario__username__icontains=search))
    try:page=int(request.GET.get('page',1))
    except ValueError:return error('Página inválida.')
    paginator=Paginator(qs,20)
    if page<1 or page>paginator.num_pages:return error('Página inexistente.',404)
    return ok(results=[{'id':e.pk,'fecha':e.fecha,'usuario':e.usuario.username if e.usuario_id else 'Sistema','recurso':e.recurso,'registro':e.registro,'accion':e.accion,'origen':e.origen,'delegacion':e.delegacion.nombre if e.delegacion_id else None,'antes':e.antes,'despues':e.despues} for e in paginator.page(page)],count=paginator.count,page=page,total_pages=paginator.num_pages)
