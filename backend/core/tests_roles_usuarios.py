from django.test import override_settings
import io
import json

import openpyxl
from django.test import TestCase, Client

from .models import Rol, Usuario

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def post_json(client, url, payload):
    return client.post(url, data=json.dumps(payload), content_type="application/json")


@override_settings(ROOT_URLCONF="core.legacy_test_urls")
class SesionBaseTestCase(TestCase):
    """Crea un administrador y un lector con sesion iniciada, y un cliente anonimo."""

    def setUp(self):
        self.rol_admin = Rol.objects.create(nombre="Administrador", descripcion="Acceso total", estado=True)
        self.rol_lector = Rol.objects.create(nombre="Lector", descripcion="Solo consulta", estado=True)
        self.admin = Usuario.objects.create_user(
            username="admin_t", password="Clave12345!", email="admin_t@laserena.cl", rol=self.rol_admin
        )
        self.lector = Usuario.objects.create_user(
            username="lector_t", password="Clave12345!", email="lector_t@laserena.cl", rol=self.rol_lector
        )
        self.client = Client()
        self.client.force_login(self.admin)
        self.client_lector = Client()
        self.client_lector.force_login(self.lector)
        self.client_anonimo = Client()


@override_settings(ROOT_URLCONF="core.legacy_test_urls")
class RolesCRUDTests(SesionBaseTestCase):
    """T09: CRUD de Roles."""

    def test_sin_sesion_devuelve_401(self):
        self.assertEqual(self.client_anonimo.get("/api/roles/").status_code, 401)

    def test_listar_roles(self):
        r = self.client.get("/api/roles/?estado=todos")
        self.assertEqual(r.status_code, 200)
        nombres = [x["nombre"] for x in r.json()["results"]]
        self.assertIn("Administrador", nombres)
        self.assertIn("Lector", nombres)

    def test_crear_rol(self):
        r = post_json(self.client, "/api/roles/", {"nombre": "Supervisor", "descripcion": "Revisa", "estado": True})
        self.assertEqual(r.status_code, 201)
        self.assertTrue(Rol.objects.filter(nombre="Supervisor").exists())

    def test_rol_lector_no_puede_crear(self):
        r = post_json(self.client_lector, "/api/roles/", {"nombre": "Intruso", "descripcion": "x", "estado": True})
        self.assertEqual(r.status_code, 403)
        self.assertFalse(Rol.objects.filter(nombre="Intruso").exists())

    def test_toggle_estado_rol(self):
        rol = Rol.objects.create(nombre="Temporal", descripcion="Sin cuentas", estado=True)
        r = self.client.post(f"/api/roles/{rol.pk}/toggle-estado/")
        self.assertEqual(r.status_code, 200)
        rol.refresh_from_db()
        self.assertFalse(rol.estado)

    def test_exportar_roles_xlsx(self):
        r = self.client.get("/api/roles/export/")
        self.assertEqual(r.status_code, 200)
        self.assertIn(XLSX_MIME, r["Content-Type"])
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        encabezados = [c.value for c in ws[1]]
        self.assertIn("Nombre", encabezados)


@override_settings(ROOT_URLCONF="core.legacy_test_urls")
class UsuariosCRUDTests(SesionBaseTestCase):
    """T10: CRUD de Usuarios."""

    def payload(self, **extra):
        data = {
            "username": "nuevo@laserena.cl",
            "email": "nuevo@laserena.cl",
            "first_name": "Nuevo",
            "last_name": "Usuario",
            "rut": None,
            "rol": self.rol_admin.pk,
            "estado": True,
            "password": "ClaveNueva123!",
        }
        data.update(extra)
        return data

    def test_sin_sesion_devuelve_401(self):
        self.assertEqual(self.client_anonimo.get("/api/usuarios/").status_code, 401)

    def test_crear_usuario_encripta_clave_y_no_la_expone(self):
        r = post_json(self.client, "/api/usuarios/", self.payload())
        self.assertEqual(r.status_code, 201)
        usuario = Usuario.objects.get(username="nuevo@laserena.cl")
        self.assertTrue(usuario.check_password("ClaveNueva123!"))
        self.assertNotEqual(usuario.password, "ClaveNueva123!")
        self.assertNotIn("password", r.json()["data"])

    def test_crear_sin_password_se_rechaza(self):
        r = post_json(self.client, "/api/usuarios/", self.payload(password=""))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["error"]["code"], "VALIDATION_ERROR")
        self.assertFalse(Usuario.objects.filter(username="nuevo@laserena.cl").exists())

    def test_email_duplicado_se_rechaza(self):
        total = Usuario.objects.count()
        r = post_json(self.client, "/api/usuarios/", self.payload(email="admin_t@laserena.cl"))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(Usuario.objects.count(), total)

    def test_username_duplicado_se_rechaza(self):
        total = Usuario.objects.count()
        r = post_json(self.client, "/api/usuarios/", self.payload(username="admin_t"))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(Usuario.objects.count(), total)

    def test_editar_sin_password_conserva_la_clave(self):
        usuario = Usuario.objects.create_user(
            username="edit@laserena.cl", password="ClaveVieja123!", email="edit@laserena.cl", first_name="Antes"
        )
        r = post_json(
            self.client,
            f"/api/usuarios/{usuario.pk}/",
            self.payload(username="edit@laserena.cl", email="edit@laserena.cl", first_name="Despues", password=""),
        )
        self.assertEqual(r.status_code, 200)
        usuario.refresh_from_db()
        self.assertEqual(usuario.first_name, "Despues")
        self.assertTrue(usuario.check_password("ClaveVieja123!"))

    def test_editar_con_password_nueva_cambia_la_clave(self):
        usuario = Usuario.objects.create_user(
            username="edit2@laserena.cl", password="ClaveVieja123!", email="edit2@laserena.cl"
        )
        r = post_json(
            self.client,
            f"/api/usuarios/{usuario.pk}/",
            self.payload(username="edit2@laserena.cl", email="edit2@laserena.cl", password="ClaveOtra456!"),
        )
        self.assertEqual(r.status_code, 200)
        usuario.refresh_from_db()
        self.assertTrue(usuario.check_password("ClaveOtra456!"))

    def test_toggle_estado_desactiva_y_reactiva(self):
        usuario = Usuario.objects.create_user(username="tog@laserena.cl", password="Clave12345!", email="tog@laserena.cl")
        self.client.post(f"/api/usuarios/{usuario.pk}/toggle-estado/")
        usuario.refresh_from_db()
        self.assertFalse(usuario.is_active)
        self.client.post(f"/api/usuarios/{usuario.pk}/toggle-estado/")
        usuario.refresh_from_db()
        self.assertTrue(usuario.is_active)

    def test_listado_no_expone_password(self):
        r = self.client.get("/api/usuarios/?page_size=50")
        self.assertEqual(r.status_code, 200)
        resultados = r.json()["results"]
        self.assertTrue(len(resultados) >= 2)
        for item in resultados:
            self.assertNotIn("password", item)

    def test_exportacion_no_incluye_contrasenas(self):
        r = self.client.get("/api/usuarios/export/")
        self.assertEqual(r.status_code, 200)
        self.assertIn(XLSX_MIME, r["Content-Type"])
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        celdas = [str(c.value) for fila in ws.iter_rows() for c in fila]
        self.assertFalse(any("contras" in c.lower() for c in celdas[: len(ws[1])]))
        self.assertFalse(any("pbkdf2" in c for c in celdas))

    def test_rol_lector_no_puede_crear_usuarios(self):
        total = Usuario.objects.count()
        r = post_json(self.client_lector, "/api/usuarios/", self.payload())
        self.assertEqual(r.status_code, 403)
        self.assertEqual(Usuario.objects.count(), total)