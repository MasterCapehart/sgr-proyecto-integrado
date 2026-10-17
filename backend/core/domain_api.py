"""Operación con permisos por registro y archivos privados; sin métricas simuladas."""
import hashlib
import json
from django.core.serializers.json import DjangoJSONEncoder
from io import BytesIO
from pathlib import Path
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db import transaction, IntegrityError
from django.db.models import Q
from django.http import FileResponse, HttpResponse
from django.utils import timezone
from django.views.decorators.http import require_http_methods
from . import local_api as api
from .models import Actividad, Evidencia, Compromiso, RevisionEvidencia, SeguimientoCompromiso, Conversacion, Mensaje, Usuario, ItemMeta

MODELS={'actividades':Actividad,'evidencias':Evidencia,'compromisos':Compromiso}
FIELDS={
 'actividades':['fecha','solicitud','accion','contacto','telefono','sector','item','compromiso','unidad','plantilla','profesional'],
 'compromisos':['fecha_solicitud','descripcion','solicitante','sector','responsable','fecha_comprometida','area_apoyo','origen','estado','unidad','codigo_verificacion_foto'],
}

def scoped(request,resource):
    qs=MODELS[resource].objects.filter(delegacion=api.active_delegation(request)).order_by('fecha_comprometida','id') if resource=='compromisos' else MODELS[resource].objects.filter(delegacion=api.active_delegation(request)).order_by('-id')
    if not api.allowed(request.user,resource,'ver_equipo') and not (resource=='evidencias' and api.allowed(request.user,'evidencias','revisar')):
        if resource=='compromisos': qs=qs.filter(Q(autor=request.user)|Q(responsable=request.user))
        elif resource=='evidencias': qs=qs.filter(actividad__autor=request.user)
        else: qs=qs.filter(autor=request.user)
    return qs

def representation(obj):
    r={'id':obj.pk,'version':obj.version,'autor':obj.autor_id,'autor_nombre':obj.autor.get_full_name() or obj.autor.username,'delegacion':obj.delegacion_id,'creado':obj.creado.isoformat(),'estado':obj.estado}
    if isinstance(obj,Actividad):
        r.update(api.serialize(obj,FIELDS['actividades']));r.update(codigo=obj.codigo,item_nombre=(obj.catalogos_snapshot or {}).get('item',{}).get('nombre',obj.item.nombre),catalogos_snapshot=obj.catalogos_snapshot,historia_catalogos='Captura posterior; etiqueta original no garantizada' if (obj.catalogos_snapshot or {}).get('observada') else 'Captura al registrar' if obj.catalogos_snapshot else 'Sin captura histórica; etiqueta actual')
        r.update(profesional_nombre=(obj.profesional.get_full_name() or obj.profesional.username) if obj.profesional_id else None,cargo=obj.cargo_id,cargo_nombre=((obj.catalogos_snapshot or {}).get('cargo') or {}).get('nombre',obj.cargo.nombre if obj.cargo_id else None),area=((obj.catalogos_snapshot or {}).get('cargo') or {}).get('area',obj.cargo.area if obj.cargo_id else None),derivado_a=obj.derivado_a,estado_derivacion=obj.estado_derivacion,ingreso_tubo='Sí' if obj.compromiso_id else 'No')
        r['derivaciones']=[{'fecha':x.fecha.isoformat(),'autor':x.autor.get_full_name() or x.autor.username,'destino':x.destino,'motivo':x.motivo,'estado':x.estado} for x in obj.derivaciones.select_related('autor')]
        r['evidencias_actuales']=[{'id':e.pk,'archivo':e.nombre_archivo,'estado':e.estado,'url':f'/api/operacion/evidencias/{e.pk}/archivo/'} for e in obj.evidencias.filter(siguiente_version__isnull=True)]
        r['imagen_verificacion']='Sí' if r['evidencias_actuales'] else 'No'
    elif isinstance(obj,Compromiso):
        r.update(api.serialize(obj,FIELDS['compromisos']));r.update(folio=obj.folio,responsable_nombre=obj.responsable.get_full_name() or obj.responsable.username,vencido=obj.fecha_comprometida<timezone.localdate() and obj.estado!='Realizado')
        from .operating_policy import current
        policy=current();r['transiciones_permitidas']=policy['transiciones'].get(obj.estado,[]);r['politica_version']=policy['version'];r['politica_validada']=policy['validada']
        r['dias_atraso_cierre']=obj.cierre_actual.dias_atraso if obj.cierre_actual_id else None
        r['cerrado_en']=obj.cierre_actual.fecha.isoformat() if obj.cierre_actual_id else None
        r['cierres']=[{'fecha':c.fecha.isoformat(),'autor':c.autor.username,'fecha_plazo':str(c.fecha_plazo),'dias_atraso':c.dias_atraso,'observacion':c.observacion} for c in obj.cierres.select_related('autor')]
        r['mes_compromiso']=obj.fecha_comprometida.strftime('%m')
        r['ultima_observacion']=obj.seguimientos.first().observacion if obj.seguimientos.exists() else ''
        r['seguimientos']=[{'fecha':x.fecha.isoformat(),'autor':x.autor.username,'estado_anterior':x.estado_anterior,'estado_nuevo':x.estado_nuevo,'observacion':x.observacion,'fecha_anterior':str(x.fecha_anterior),'fecha_nueva':str(x.fecha_nueva),'responsable_anterior':x.responsable_anterior.username,'responsable_nuevo':x.responsable_nuevo.username} for x in obj.seguimientos.select_related('autor','responsable_anterior','responsable_nuevo')]
    else:
        r.update(actividad=obj.actividad_id,codigo_actividad=obj.actividad.codigo,descripcion=obj.descripcion,nombre_archivo=obj.nombre_archivo,tamano=obj.tamano,sha256=obj.sha256,reemplaza=obj.reemplaza_id,vigente=not hasattr(obj,'siguiente_version'),archivo_url=f'/api/operacion/evidencias/{obj.pk}/archivo/')
        r['revisiones']=[{'fecha':x.fecha.isoformat(),'autor':x.autor.username,'decision':x.decision,'observacion':x.observacion,'version_evidencia':x.version_evidencia} for x in obj.revisiones.select_related('autor')]
    return json.loads(json.dumps(r,cls=DjangoJSONEncoder))

def validate_version(obj,data):
    if '_version' not in data:return api.error('Falta la versión del registro.',428)
    if str(data['_version'])!=str(obj.version):return api.error('El registro cambió. Recarga antes de guardar.',409)

def apply_fields(request,resource,obj,data):
    if set(data)-set(FIELDS[resource])-{'_version','observacion'}:raise ValidationError('Campos no permitidos.')
    for field in FIELDS[resource]:
        if field not in data:continue
        value=data[field]
        if field in ['item','compromiso','responsable','profesional']:
            if field in ['compromiso','profesional'] and value in ('',None):setattr(obj,field,None);continue
            if isinstance(value,bool) or not str(value).isdigit():raise ValidationError('Referencia inválida.')
            if field=='item':
                from .excel_support import items_for
                related=items_for(request.user).filter(pk=value).first()
            elif field=='compromiso': related=scoped(request,'compromisos').filter(pk=value).first()
            else: related=Usuario.objects.filter(pk=value,is_active=True).filter(Q(delegacion=obj.delegacion)|Q(delegaciones=obj.delegacion)).distinct().first()
            if related is None:raise ValidationError(f'{field}: registro inactivo o fuera de la delegación.')
            setattr(obj,field,related)
        elif field=='unidad':
            if isinstance(value,bool) or str(value)!='1':raise ValidationError('La unidad de la matriz es 1 por registro.')
            obj.unidad=int(value)
        elif field=='plantilla':
            from .matrix_profiles import PROFILES
            selected=request.user.cargo.planilla if request.user.cargo_id else 'GENERAL'
            if obj.pk:
                if value!=obj.plantilla:raise ValidationError('La planilla del registro original no se cambia.')
            elif value not in PROFILES or (not request.user.is_superuser and value!=(selected or 'GENERAL')):raise ValidationError('Planilla no autorizada para tu función.')
            obj.plantilla=value
        else:
            if field=='codigo_verificacion_foto' and value in ['',None]:obj.codigo_verificacion_foto='';continue
            if field=='fecha_solicitud' and value in ('',None):obj.fecha_solicitud=None;continue
            if not isinstance(value,str) or not value.strip():raise ValidationError(f'{field}: campo obligatorio de texto.')
            setattr(obj,field,value.strip())
    if resource=='actividades':
        import re
        if obj.telefono and not re.fullmatch(r'\+?[0-9 ()-]{8,20}',obj.telefono):raise ValidationError('Teléfono inválido.')
        from datetime import date
        try: day=date.fromisoformat(str(obj.fecha))
        except ValueError:raise ValidationError('Fecha inválida.')
        if day>timezone.localdate():raise ValidationError('La actividad no puede tener fecha futura.')
    if resource=='compromisos' and obj.fecha_solicitud:
        from datetime import date
        try: requested=date.fromisoformat(str(obj.fecha_solicitud));promised=date.fromisoformat(str(obj.fecha_comprometida))
        except ValueError:raise ValidationError('Fecha de solicitud o compromiso inválida.')
        if requested>promised:raise ValidationError('La fecha de compromiso no puede preceder a la solicitud.')
    obj.full_clean()

def gate(request,resource,action):
    if not request.user.is_authenticated:return api.error('Inicia sesión.',401)
    if not api.allowed(request.user,resource,action):return api.error('No tienes permiso para esta operación.',403)
    if api.active_delegation(request) is None:return api.error('Selecciona una delegación autorizada.',403)

@require_http_methods(['GET'])
def options(request):
    denied=gate(request,'actividades','ver')
    if denied and not api.allowed(request.user,'compromisos','ver') and not api.allowed(request.user,'evidencias','ver'):return denied
    d=api.active_delegation(request)
    from .excel_support import items_for
    from .models import Cargo
    return api.ok(plantilla=(request.user.cargo.planilla if request.user.cargo_id else '') or 'GENERAL',cargos=list(Cargo.objects.filter(estado=True).values('id','nombre','area')),items=list(items_for(request.user).values('id','codigo','nombre')),responsables=[{'id':u.pk,'nombre':u.get_full_name() or u.username} for u in Usuario.objects.filter(is_active=True).filter(Q(delegacion=d)|Q(delegaciones=d)).distinct()])

@require_http_methods(['GET','POST','PATCH'])
def records(request,resource,pk=None,action=None):
    permission='exportar' if action=='export' else ('crear' if request.method=='POST' else 'editar' if request.method=='PATCH' else 'ver')
    denied=gate(request,resource,permission)
    if denied:return denied
    if action=='export' and request.method!='GET':return api.error('Método no permitido.',405)
    try:
        qs=scoped(request,resource)
        if request.method=='GET':
            if pk:
                obj=qs.filter(pk=pk).first()
                return api.ok(representation(obj)) if obj else api.error('Registro no encontrado.',404)
            search=request.GET.get('search','').strip()
            if search:
                query=Q(descripcion__icontains=search) if resource!='actividades' else Q(solicitud__icontains=search)|Q(accion__icontains=search)
                code='codigo' if resource=='actividades' else 'folio' if resource=='compromisos' else 'nombre_archivo'
                qs=qs.filter(query|Q(**{code+'__icontains':search}))
            from .excel_support import COLUMN_FILTERS
            for key,value in request.GET.items():
                if key.startswith('col_') and value:
                    lookup=COLUMN_FILTERS[resource].get(key[4:])
                    if lookup is None:return api.error('Filtro de columna inválido.')
                    if key in ['col_autor_nombre','col_responsable_nombre']:
                        from django.db.models.functions import Concat
                        from django.db.models import Value, CharField
                        relation='autor' if key=='col_autor_nombre' else 'responsable'
                        qs=qs.annotate(_nombre_filtro=Concat(relation+'__first_name',Value(' '),relation+'__last_name',output_field=CharField())).filter(Q(_nombre_filtro__icontains=value)|Q(**{relation+'__username__icontains':value}))
                    else:qs=qs.filter(**{lookup:value})
            if resource=='actividades' and request.GET.get('plantilla'):qs=qs.filter(plantilla=request.GET['plantilla'])
            if request.GET.get('estado'):qs=qs.filter(estado=request.GET['estado'])
            if request.GET.get('sector') and resource!='evidencias':qs=qs.filter(sector__icontains=request.GET['sector'])
            if resource=='compromisos' and request.GET.get('plazo')=='vencido':qs=qs.filter(fecha_comprometida__lt=timezone.localdate()).exclude(estado='Realizado')
            if resource=='compromisos' and request.GET.get('plazo')=='hoy':qs=qs.filter(fecha_comprometida=timezone.localdate()).exclude(estado='Realizado')
            for key in ['autor','item','actividad','responsable','cargo']:
                if request.GET.get(key) and key in [f.name for f in MODELS[resource]._meta.fields]:qs=qs.filter(**{key:request.GET[key]})
            datefield='fecha' if resource=='actividades' else 'fecha_comprometida' if resource=='compromisos' else 'creado__date'
            for key,op in [('desde','gte'),('hasta','lte')]:
                if request.GET.get(key):qs=qs.filter(**{datefield+'__'+op:request.GET[key]})
            if resource=='actividades' and request.GET.get('ingreso_tubo') in ['si','no']:qs=qs.filter(compromiso__isnull=request.GET['ingreso_tubo']=='no')
            if resource=='actividades' and request.GET.get('evidencia') in ['si','no']:
                evidence_ids=Evidencia.objects.filter(siguiente_version__isnull=True).values('actividad_id')
                qs=qs.filter(pk__in=evidence_ids) if request.GET['evidencia']=='si' else qs.exclude(pk__in=evidence_ids)
            if action=='export':
                from openpyxl import Workbook
                book=Workbook();sheet=book.active;sheet.title='Registros SGR';sheet.append(['Recurso',resource]);sheet.append(['Generado',timezone.localtime().isoformat()]);sheet.append(['Delegación',api.active_delegation(request).nombre])
                from .excel_support import EXCEL_COLUMNS
                from openpyxl.styles import Font, PatternFill, Alignment
                columns=EXCEL_COLUMNS[resource]
                if resource=='actividades' and request.GET.get('vista'):
                    from .matrix_profiles import PROFILES
                    code=request.GET['vista']
                    if code not in PROFILES:raise ValidationError('Vista de planilla inválida.')
                    if PROFILES[code].get('ruta'):raise ValidationError('Exporta esta planilla desde Gestión social.')
                    columns=PROFILES[code]['columnas']
                sheet.append([title for _,title in columns])
                for obj in qs:
                    r=representation(obj);sheet.append([str(r.get(f) or '') for f,_ in columns])
                    for cell in sheet[sheet.max_row]:cell.data_type='s';cell.alignment=Alignment(vertical='top',wrap_text=True)
                for cell in sheet[4]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='AD0000')
                sheet.freeze_panes='C5';sheet.auto_filter.ref=f'A4:{sheet.cell(sheet.max_row,len(columns)).coordinate}'
                for i,(key,title) in enumerate(columns,1):sheet.column_dimensions[sheet.cell(4,i).column_letter].width=50 if key in ['descripcion','solicitud','accion','ultima_observacion'] else 24
                out=BytesIO();book.save(out);response=HttpResponse(out.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet');response['Content-Disposition']=f'attachment; filename="{resource}.xlsx"';api.record_event(request,resource,None,'exportar',None,{'cantidad':qs.count()});return response
            page=int(request.GET.get('page',1));size=int(request.GET.get('page_size',20))
            if page<1 or not 1<=size<=100:return api.error('Paginación inválida.')
            pager=Paginator(qs,size)
            if page>pager.num_pages:return api.error('Página inexistente.',404)
            from .kpis import resource_kpis
            return api.ok(results=[representation(o) for o in pager.page(page)],count=pager.count,page=page,total_pages=pager.num_pages,kpis=resource_kpis(qs,resource))
        if resource=='evidencias':
            if request.method!='POST' or pk:return api.error('Las evidencias se conservan; adjunta una nueva versión.',405)
            return upload(request)
        data=api.body(request)
        with transaction.atomic():
            if request.method=='PATCH':
                obj=qs.select_for_update().filter(pk=pk).first()
                if not obj:return api.error('Registro no encontrado.',404)
                if obj.autor_id!=request.user.pk and not (resource=='compromisos' and obj.responsable_id==request.user.pk) and not api.allowed(request.user,resource,'editar_equipo'):return api.error('No puedes editar registros de otro responsable.',403)
                if resource=='actividades' and obj.estado not in ['Borrador','Corrección','Rechazada']:return api.error('Solo se editan actividades en borrador o devueltas.')
                conflict=validate_version(obj,data)
                if conflict:return conflict
                before=representation(obj)
            else:
                if pk:return api.error('Ruta inválida.',405)
                obj=MODELS[resource](autor=request.user,delegacion=api.active_delegation(request));before=None
                if resource=='actividades':obj.cargo=request.user.cargo;obj.plantilla=(request.user.cargo.planilla if request.user.cargo_id else '') or 'GENERAL'
                if resource=='compromisos' and not data.get('fecha_solicitud'):obj.fecha_solicitud=timezone.localdate()
            old_item=obj.item_id if resource=='actividades' else None
            apply_fields(request,resource,obj,data)
            if resource=='actividades' and (not obj.catalogos_snapshot or old_item!=obj.item_id):
                from .catalog_history import activity
                obj.catalogos_snapshot=activity(obj,observed=bool(before and old_item==obj.item_id))
            if resource=='compromisos':
                from .closure_workflow import validate_transition
                validate_transition(request,obj,before,data)
            if before:obj.version+=1
            obj.save()
            if resource=='compromisos' and before:
                observation=data.get('observacion')
                if not isinstance(observation,str) or not observation.strip() or len(observation)>2000:raise ValidationError('Indica una observación de hasta 2000 caracteres.')
                observation=observation.strip()
                SeguimientoCompromiso.objects.create(compromiso=obj,autor=request.user,estado_anterior=before['estado'],estado_nuevo=obj.estado,observacion=observation,responsable_anterior_id=before['responsable'],responsable_nuevo=obj.responsable,fecha_anterior=before['fecha_comprometida'],fecha_nueva=obj.fecha_comprometida)
            if resource=='compromisos':
                from .closure_workflow import save_closure
                save_closure(request,obj,before,data)
            api.record_event(request,resource,obj.pk,'editar' if before else 'crear',before,representation(obj))
        response=api.ok(representation(obj));response.status_code=200 if before else 201;return response
    except (ValidationError,ValueError,TypeError) as ex:return api.error('; '.join(ex.messages) if isinstance(ex,ValidationError) else 'Datos inválidos.')
    except IntegrityError:return api.error('El registro cambió o existe una versión posterior.',409)

def upload(request):
    data=request.POST;file=request.FILES.get('archivo');activity=scoped(request,'actividades').filter(pk=data.get('actividad') if str(data.get('actividad','')).isdigit() else 0).first()
    if not activity:return api.error('Actividad no encontrada.',404)
    if not api.allowed(request.user,'actividades','editar') or (activity.autor_id!=request.user.pk and not api.allowed(request.user,'actividades','editar_equipo')):return api.error('No puedes adjuntar a esta actividad.',403)
    if activity.estado=='Validada':return api.error('La actividad está validada; conserva el respaldo.')
    if not file or file.size>10*1024*1024 or file.size==0:return api.error('Adjunta un archivo de hasta 10 MB.')
    signature=file.read(12);file.seek(0);ext=Path(file.name).suffix.lower()
    mime='application/pdf' if ext=='.pdf' and signature.startswith(b'%PDF-') else 'image/png' if ext=='.png' and signature.startswith(b'\x89PNG\r\n\x1a\n') else 'image/jpeg' if ext in ['.jpg','.jpeg'] and signature.startswith(b'\xff\xd8\xff') else None
    if not mime:return api.error('Solo se admiten PDF, PNG y JPEG con contenido compatible.')
    digest=hashlib.sha256()
    for chunk in file.chunks():digest.update(chunk)
    file.seek(0);saved=None
    try:
        with transaction.atomic():
            activity=Actividad.objects.select_for_update().get(pk=activity.pk)
            if activity.estado=='Validada':raise ValidationError('La actividad ya está validada.')
            previous=None
            if data.get('reemplaza'):
                previous=Evidencia.objects.select_for_update().filter(pk=data['reemplaza'],actividad=activity).first()
                if not previous or hasattr(previous,'siguiente_version') or previous.estado not in ['Corrección','Rechazada','Anulada']:raise ValidationError('Solo puedes reemplazar un respaldo devuelto y vigente.')
            obj=Evidencia(actividad=activity,autor=request.user,delegacion=activity.delegacion,archivo=file,nombre_archivo=Path(file.name).name,tipo_mime=mime,tamano=file.size,sha256=digest.hexdigest(),descripcion=data.get('descripcion','').strip(),reemplaza=previous)
            obj.full_clean();obj.save();saved=obj.archivo
            activity.estado='Pendiente';activity.version+=1;activity.save()
            api.record_event(request,'evidencias',obj.pk,'adjuntar',None,representation(obj))
        response=api.ok(representation(obj));response.status_code=201;return response
    except Exception:
        if saved:saved.delete(save=False)
        raise

@require_http_methods(['GET'])
def download(request,pk):
    denied=gate(request,'evidencias','ver')
    if denied:return denied
    obj=scoped(request,'evidencias').filter(pk=pk).first()
    if not obj:return api.error('Archivo no encontrado.',404)
    try:response=FileResponse(obj.archivo.open('rb'),as_attachment=True,filename=obj.nombre_archivo,content_type=obj.tipo_mime)
    except FileNotFoundError:return api.error('El archivo no está disponible.',404)
    response['X-Content-Type-Options']='nosniff';response['Cache-Control']='private, no-store';return response

@require_http_methods(['POST'])
def review(request,pk):
    denied=gate(request,'evidencias','revisar')
    if denied:return denied
    try:
        data=api.body(request)
        with transaction.atomic():
            candidate=scoped(request,'evidencias').filter(pk=pk).values('actividad_id').first()
            if not candidate:return api.error('Evidencia no encontrada.',404)
            activity=Actividad.objects.select_for_update().get(pk=candidate['actividad_id'])
            obj=scoped(request,'evidencias').select_for_update().filter(pk=pk).first()
            if not obj:return api.error('Evidencia no encontrada.',404)
            if request.user.pk in [obj.autor_id,obj.actividad.autor_id]:return api.error('El autor no puede revisar su propia evidencia.')
            conflict=validate_version(obj,data)
            if conflict:return conflict
            if hasattr(obj,'siguiente_version') or obj.estado!='Pendiente':return api.error('Esta versión no está pendiente.')
            if data.get('decision') not in ['Aprobada','Rechazada','Corrección']:return api.error('Decisión inválida.')
            observation=data.get('observacion')
            if not isinstance(observation,str) or not observation.strip() or len(observation)>2000:return api.error('Indica una observación de hasta 2000 caracteres.')
            before=representation(obj);obj.estado=data['decision'];obj.version+=1;obj.save()
            RevisionEvidencia.objects.create(evidencia=obj,autor=request.user,decision=obj.estado,observacion=observation.strip(),version_evidencia=obj.version)
            activity=Actividad.objects.select_for_update().get(pk=obj.actividad_id)
            current=activity.evidencias.filter(siguiente_version__isnull=True)
            states=set(current.exclude(estado='Anulada').values_list('estado',flat=True));activity.estado='Validada' if states=={'Aprobada'} else 'Corrección' if 'Corrección' in states else 'Rechazada' if 'Rechazada' in states else 'Pendiente';activity.version+=1;activity.save()
            api.record_event(request,'evidencias',obj.pk,'revisar',before,representation(obj))
            api.record_event(request,'actividades',activity.pk,'validar',None,representation(activity))
        return api.ok(representation(obj))
    except (ValidationError,ValueError,TypeError):return api.error('Datos de revisión inválidos.')

@require_http_methods(['GET','POST'])
def messages(request,resource,pk):
    denied=gate(request,resource,'ver')
    if denied:return denied
    obj=scoped(request,resource).filter(pk=pk).first()
    if not obj:return api.error('Contexto no encontrado.',404)
    context={'actividades':'actividad','compromisos':'compromiso','evidencias':'evidencia'}[resource]
    if request.method=='POST':
        if not api.allowed(request.user,'conversaciones','crear'):return api.error('No puedes escribir comentarios.',403)
        try:data=api.body(request)
        except ValidationError:return api.error('Mensaje inválido.')
        content=data.get('contenido')
        if not isinstance(content,str) or not content.strip() or len(content)>2000:return api.error('Escribe un mensaje de hasta 2000 caracteres.')
        with transaction.atomic():
            conversation,_=Conversacion.objects.get_or_create(**{context:obj})
            msg=Mensaje.objects.create(conversacion=conversation,autor=request.user,contenido=content.strip())
            api.record_event(request,'conversaciones',conversation.pk,'comentar',None,{'mensaje':msg.pk,'contexto':resource,'registro':pk})
        response=api.ok({'id':msg.pk});response.status_code=201;return response
    conversation=Conversacion.objects.filter(**{context:obj}).first()
    return api.ok(results=[{'id':m.pk,'contenido':m.contenido,'autor':m.autor.get_full_name() or m.autor.username,'fecha':m.fecha.isoformat()} for m in conversation.mensajes.select_related('autor')] if conversation else [])
