from django.test import TestCase
from .models import Usuario

class LimitesLocales(TestCase):
    def test_anonimo_no_lee_catalogos(self):
        self.assertEqual(self.client.get('/api/tipos-atencion/').status_code, 401)

    def test_roles_disponibles_para_admin(self):
        user = Usuario.objects.create_superuser('admin_test', 'admin@example.test', 'EjemploFuerte!247')
        self.client.force_login(user)
        self.assertEqual(self.client.get('/api/roles/').status_code, 200)

    def test_no_admite_json_lista(self):
        user = Usuario.objects.create_superuser('admin_test', 'admin@example.test', 'EjemploFuerte!247')
        self.client.force_login(user)
        self.assertEqual(self.client.post('/api/tipos-atencion/', '[]', content_type='application/json').status_code, 400)

import re
from io import BytesIO
from datetime import timedelta
from django.test import override_settings
from .test_support import Client
from django.core import mail
from django.utils import timezone
from openpyxl import load_workbook
from .models import Delegacion, Rol, TipoAtencion, SubAtencion, ItemMeta, ConfiguracionMeta, Atencion, RecuperacionAcceso
from .local_api import PERMISSIONS

@override_settings(EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
class FlujoLocal(TestCase):
    def setUp(self):
        self.d=Delegacion.objects.create(codigo='D1',nombre='Delegación Uno')
        self.other=Delegacion.objects.create(codigo='D2',nombre='Delegación Dos')
        self.role=Rol.objects.create(nombre='Operador',permisos=['atenciones.ver','atenciones.crear','atenciones.exportar','tipos-atencion.ver','subatenciones.ver'])
        self.user=Usuario.objects.create_user('operador','operador@example.test','ClaveInicial!456',rol=self.role,delegacion=self.d)
        self.user.delegaciones.add(self.d)
        self.admin=Usuario.objects.create_superuser('administrador','admin@example.test','ClaveAdmin!456')
        self.client=Client()
        self.client.force_login(self.admin)
        self.tipo=TipoAtencion.objects.create(codigo='T1',nombre='Consulta')
        self.sub=SubAtencion.objects.create(codigo='S1',nombre='Orientación',tipo_atencion=self.tipo)

    def send(self,path,data,method='post'):
        import json
        return getattr(self.client,method)(path,json.dumps(data),content_type='application/json')

    def attention(self,**changes):
        return {'solicitante':'Vecina Ficticia','tipo_atencion':self.tipo.pk,'sub_atencion':self.sub.pk,'descripcion':'Consulta de prueba',**changes}

    def test_login_correo_y_logout(self):
        self.client.logout()
        self.assertEqual(self.send('/api/login/',{'username':self.user.email,'password':'ClaveInicial!456'}).status_code,200)
        self.assertEqual(self.client.get('/api/contexto/').json()['data']['usuario']['id'],self.user.pk)
        self.assertEqual(self.send('/api/logout/',{}).status_code,200)
        self.assertEqual(self.client.get('/api/contexto/').status_code,401)

    def test_login_incorrecto_inactivo_y_limite(self):
        self.client.logout()
        self.user.is_active=False;self.user.save()
        self.assertEqual(self.send('/api/login/',{'username':self.user.username,'password':'ClaveInicial!456'}).status_code,401)
        for _ in range(4):self.assertEqual(self.send('/api/login/',{'username':'inexistente','password':'mal'}).status_code,401)
        self.assertEqual(self.send('/api/login/',{}).status_code,429)

    def test_csrf_requerido(self):
        c=Client(enforce_csrf_checks=True);c.force_login(self.admin)
        self.assertEqual(c.post('/api/roles/', '{}',content_type='application/json').status_code,403)

    def test_permiso_y_delegacion(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.get('/api/roles/').status_code,403)
        self.assertEqual(self.send('/api/contexto/cambiar-delegacion/',{'delegacion_id':self.other.pk}).status_code,403)
        self.assertEqual(self.send('/api/contexto/cambiar-delegacion/',{'delegacion_id':self.d.pk}).status_code,200)
        self.user.refresh_from_db();self.assertEqual(self.user.delegacion_id,self.d.pk)
        self.assertEqual(self.send('/api/atenciones/',self.attention(delegacion_responsable=self.other.pk)).status_code,403)

    def test_atencion_par_incorrecto_y_doble_envio(self):
        self.client.force_login(self.user)
        other=TipoAtencion.objects.create(codigo='OTRO',nombre='Otro')
        self.assertEqual(self.send('/api/atenciones/',self.attention(tipo_atencion=other.pk)).status_code,400)
        payload=self.attention(clave_idempotencia='unico-local')
        first=self.send('/api/atenciones/',payload);self.assertEqual(first.status_code,201,first.content)
        second=self.send('/api/atenciones/',payload);self.assertEqual(second.status_code,200)
        self.assertEqual(Atencion.objects.count(),1)
        self.assertEqual(first.json()['data']['usuario'],self.user.pk)
        self.assertEqual(self.send('/api/atenciones/',self.attention(clave_idempotencia='unico-local',descripcion='Diferente')).status_code,409)

    def test_atenciones_aisladas_en_lista_detalle_y_exportacion(self):
        remote=Atencion.objects.create(folio='REMOTO',solicitante='Otra Persona',tipo_atencion=self.tipo,sub_atencion=self.sub,descripcion='Privado',delegacion=self.other,usuario=self.admin)
        self.client.force_login(self.user)
        self.assertEqual(self.client.get('/api/atenciones/').json()['count'],0)
        self.assertEqual(self.client.get(f'/api/atenciones/{remote.pk}/').status_code,404)
        response=self.client.get('/api/atenciones/export/')
        sheet=load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row,1)

    def test_crud_seis_recursos(self):
        specs={
          'delegaciones':{'codigo':'D3','nombre':'Delegación Tres','direccion':'Calle ficticia 10','comuna':'La Serena'},
          'roles':{'nombre':'Consulta local','descripcion':'Solo consulta','permisos':['tipos-atencion.ver'],'alcance':'delegacion'},
          'usuarios':{'username':'nuevo','first_name':'Persona','last_name':'Prueba','email':'nuevo@example.test','password':'Distinta!897342','rol':self.role.pk,'delegacion':self.d.pk,'delegaciones':[self.d.pk]},
          'items-metas':{'codigo':'I1','nombre':'Ítem Uno','descripcion':'Prueba'},
          'tipos-atencion':{'codigo':'T2','nombre':'Tipo Dos','descripcion':'Prueba'},
          'subatenciones':{'codigo':'S2','nombre':'Sub Dos','tipo_atencion':self.tipo.pk},
        }
        for resource,payload in specs.items():
            with self.subTest(resource=resource):
                url=f'/api/{resource}/';r=self.send(url,payload);self.assertEqual(r.status_code,201,r.content);pk=r.json()['data']['id']
                self.assertEqual(self.client.get(f'{url}{pk}/').status_code,200)
                edit={'first_name':'Nombre Editado'} if resource=='usuarios' else {'nombre':'Registro Editado '+resource}
                self.assertEqual(self.send(f'{url}{pk}/',edit,'patch').status_code,200)
                self.assertEqual(self.client.delete(f'{url}{pk}/',content_type='application/json').status_code,200)
                state=self.client.get(f'{url}{pk}/').json()['data'];self.assertFalse(state.get('is_active',state.get('estado')))

    def test_duplicado_y_longitud(self):
        self.assertEqual(self.send('/api/tipos-atencion/',{'codigo':'t1','nombre':'Otro'}).status_code,400)
        self.assertEqual(self.send('/api/tipos-atencion/',{'codigo':'N'*21,'nombre':'Otro'}).status_code,400)
        self.assertEqual(self.send('/api/tipos-atencion/',{'codigo':'N','nombre':''}).status_code,400)

    def test_no_bypass_relaciones_por_patch(self):
        self.assertEqual(self.send(f'/api/tipos-atencion/{self.tipo.pk}/',{'estado':False},'patch').status_code,400)
        self.assertEqual(self.client.delete(f'/api/delegaciones/{self.d.pk}/').status_code,400)
        self.assertEqual(self.client.delete(f'/api/roles/{self.role.pk}/').status_code,400)
        self.assertEqual(self.send(f'/api/subatenciones/{self.sub.pk}/',{'tipo_atencion':99999},'patch').status_code,400)

    def test_usuarios_sin_secretos_y_sin_elevacion(self):
        response=self.client.get('/api/usuarios/').json()
        self.assertNotIn('password',str(response));self.assertNotIn('codigo_recuperacion',str(response))
        self.assertEqual(self.send(f'/api/usuarios/{self.user.pk}/',{'is_superuser':True},'patch').status_code,400)
        self.assertEqual(self.send(f'/api/usuarios/{self.user.pk}/',{'email':self.admin.email},'patch').status_code,400)
        self.role.permisos+=['roles.crear'];self.role.save();self.client.force_login(self.user)
        self.assertEqual(self.send('/api/roles/',{'nombre':'Elevar','permisos':PERMISSIONS,'alcance':'global'}).status_code,400)

    def test_paginacion_filtros_exporta_todas_las_paginas(self):
        for i in range(13):ItemMeta.objects.create(codigo=f'I{i}',nombre=f'Coincidencia {i}')
        response=self.client.get('/api/items-metas/?search=COINCIDENCIA&page_size=5&page=2').json()
        self.assertEqual(response['count'],13);self.assertEqual(len(response['results']),5);self.assertEqual(response['page'],2)
        excel=self.client.get('/api/items-metas/export/?search=coincidencia&page_size=5&page=2')
        sheet=load_workbook(BytesIO(excel.content)).active;self.assertEqual(sheet.max_row,14)
        self.assertEqual(self.client.get('/api/items-metas/?ordering=password').status_code,400)
        self.assertEqual(self.client.get('/api/items-metas/?page_size=0').status_code,400)

    def test_excel_no_ejecuta_formulas(self):
        ItemMeta.objects.create(codigo='FORM',nombre='=1+2')
        sheet=load_workbook(BytesIO(self.client.get('/api/items-metas/export/').content)).active
        self.assertEqual(sheet.cell(2,3).data_type,'s')

    def request_recovery(self):
        self.client.logout();r=self.send('/api/recuperacion/',{'email':self.user.email});self.assertEqual(r.status_code,200)
        return re.search(r'\b\d{6}\b',mail.outbox[-1].body).group()

    def test_recuperacion_completa_y_un_solo_uso(self):
        code=self.request_recovery();rec=RecuperacionAcceso.objects.get();self.assertNotEqual(rec.codigo_hash,code)
        self.assertEqual(self.send('/api/recuperacion/cambiar/',{'password':'NuevaClaveFuerte!789'}).status_code,403)
        self.assertEqual(self.send('/api/recuperacion/verificar/',{'codigo':'incorrecto'}).status_code,400)
        self.assertEqual(self.send('/api/recuperacion/verificar/',{'codigo':code}).status_code,200)
        self.assertEqual(self.send('/api/recuperacion/cambiar/',{'password':'123','confirmacion':'123'}).status_code,400)
        self.assertEqual(self.send('/api/recuperacion/cambiar/',{'password':'NuevaClaveFuerte!789','confirmacion':'Otra'}).status_code,400)
        self.assertEqual(self.send('/api/recuperacion/cambiar/',{'password':'NuevaClaveFuerte!789','confirmacion':'NuevaClaveFuerte!789'}).status_code,200)
        self.assertEqual(self.send('/api/recuperacion/verificar/',{'codigo':code}).status_code,400)
        self.assertEqual(self.send('/api/login/',{'username':self.user.username,'password':'ClaveInicial!456'}).status_code,401)
        self.assertEqual(self.send('/api/login/',{'username':self.user.username,'password':'NuevaClaveFuerte!789'}).status_code,200)

    def test_codigo_vencido_y_reenvio(self):
        code=self.request_recovery()
        self.assertEqual(self.send('/api/recuperacion/',{'email':self.user.email}).status_code,429)
        RecuperacionAcceso.objects.update(vence=timezone.now()-timedelta(seconds=1))
        self.assertEqual(self.send('/api/recuperacion/verificar/',{'codigo':code}).status_code,400)

    def test_intentos_codigo(self):
        code=self.request_recovery()
        for _ in range(5):self.assertEqual(self.send('/api/recuperacion/verificar/',{'codigo':'no'}).status_code,400)
        self.assertEqual(self.send('/api/recuperacion/verificar/',{'codigo':code}).status_code,400)

    def test_sin_enumeracion_correo(self):
        self.client.logout();r=self.send('/api/recuperacion/',{'email':'no-existe@example.test'});self.assertEqual(r.status_code,200)
        self.assertFalse(RecuperacionAcceso.objects.exists())
        self.assertEqual(self.send('/api/recuperacion/verificar/',{'codigo':'123456'}).status_code,400)

    def test_caso_maximo_tres_y_misma_persona(self):
        self.client.force_login(self.user)
        for i in range(3):self.assertEqual(self.send('/api/atenciones/',self.attention(caso_social='CS-1')).status_code,201)
        self.assertEqual(self.send('/api/atenciones/',self.attention(caso_social='CS-1')).status_code,409)
        self.assertEqual(self.send('/api/atenciones/',self.attention(caso_social='CS-1',solicitante='Otra Persona')).status_code,409)

    def test_post_detalle_no_elude_permiso_edicion(self):
        self.assertEqual(self.send(f'/api/roles/{self.role.pk}/',{'nombre':'No permitido'}).status_code,405)
        self.assertEqual(self.client.get(f'/api/roles/{self.role.pk}/toggle-estado/').status_code,405)

    def test_export_usuarios_no_incluye_secretos(self):
        sheet=load_workbook(BytesIO(self.client.get('/api/usuarios/export/').content)).active
        headers=list(next(sheet.values))
        for secret in ['password','codigo_recuperacion','correo_unico','last_login']: self.assertNotIn(secret,headers)

    def test_modificar_rol_superior_rechazado(self):
        elevated=Rol.objects.create(nombre='Superior',permisos=PERMISSIONS)
        self.role.permisos+=['roles.editar'];self.role.save();self.client.force_login(self.user)
        self.assertEqual(self.send(f'/api/roles/{elevated.pk}/',{'permisos':[]},'patch').status_code,403)

    def test_idempotencia_no_acepta_otro_tipo(self):
        self.client.force_login(self.user)
        p=self.attention(clave_idempotencia='clave-tipo');self.assertEqual(self.send('/api/atenciones/',p).status_code,201)
        other=TipoAtencion.objects.create(codigo='DIFF',nombre='Distinto')
        self.assertEqual(self.send('/api/atenciones/',{**p,'tipo_atencion':other.pk}).status_code,409)

    def test_meta_validada_y_relacion_protegida(self):
        item=ItemMeta.objects.create(codigo='META',nombre='Ítem Meta')
        self.assertEqual(self.send('/api/configuraciones-metas/',{'item':item.pk,'periodo':'2026','meta_esperada':-1,'ponderacion':20}).status_code,400)
        self.assertEqual(self.send('/api/configuraciones-metas/',{'item':item.pk,'periodo':'2026','meta_esperada':10,'ponderacion':20}).status_code,201)
        self.assertEqual(self.client.delete(f'/api/items-metas/{item.pk}/').status_code,400)

    def test_cambio_contexto_no_modifica_delegacion_principal(self):
        self.user.delegaciones.add(self.other);self.client.force_login(self.user)
        self.assertEqual(self.send('/api/contexto/cambiar-delegacion/',{'delegacion_id':self.other.pk}).status_code,200)
        self.user.refresh_from_db();self.assertEqual(self.user.delegacion_id,self.d.pk)
        r=self.send('/api/atenciones/',self.attention());self.assertEqual(r.status_code,201);self.assertEqual(r.json()['data']['delegacion'],self.other.pk)

    def test_reenvio_invalida_codigo_anterior(self):
        from unittest.mock import patch
        old=self.request_recovery()
        session=self.client.session;session['recovery_sent']=0;session.save()
        with patch('core.local_api.timezone.now',return_value=timezone.now()+timedelta(seconds=61)):
            self.assertEqual(self.send('/api/recuperacion/',{'email':self.user.email}).status_code,200)
        recs=RecuperacionAcceso.objects.order_by('pk');self.assertTrue(recs[0].consumido);self.assertFalse(recs[1].consumido)
        # La sesión ahora apunta al desafío nuevo; usar el anterior no debe validarlo.
        recs[1].codigo_hash='pbkdf2_sha256$1000000$distinto$incorrecto';recs[1].save()
        self.assertEqual(self.send('/api/recuperacion/verificar/',{'codigo':old}).status_code,400)

    def test_nueva_sesion_no_reinicia_limite_login(self):
        for _ in range(5):
            c=Client();self.assertEqual(c.post('/api/login/',{'username':'ausente','password':'incorrecta'},content_type='application/json').status_code,401)
        c=Client();self.assertEqual(c.post('/api/login/',{'username':'ausente','password':'incorrecta'},content_type='application/json').status_code,429)

    def test_estado_recuperacion_exige_verificacion(self):
        self.assertEqual(self.client.get('/api/recuperacion/estado/').status_code,403)
        code=self.request_recovery();self.send('/api/recuperacion/verificar/',{'codigo':code})
        self.assertEqual(self.client.get('/api/recuperacion/estado/').status_code,200)

    def test_formatos_no_textuales_rechazados(self):
        self.assertEqual(self.send('/api/tipos-atencion/',{'codigo':'INVALID','nombre':{'inesperado':True}}).status_code,400)
        self.assertEqual(self.send('/api/atenciones/',self.attention(solicitante=['Persona'])).status_code,400)

class ValidacionesPorMaestro(TestCase):
    setUp=FlujoLocal.setUp
    send=FlujoLocal.send

SPECIFIC={
 'delegaciones':({'codigo':'NUEVA','nombre':'Nueva Delegación','direccion':'Calle ficticia','comuna':'La Serena'},'nombre',100,['codigo','nombre']),
 'roles':({'nombre':'Nuevo Rol','descripcion':'Prueba','alcance':'delegacion','permisos':[]},'nombre',50,['nombre']),
 'usuarios':({'username':'nuevo-campos','first_name':'Nombre','last_name':'Apellido','email':'campos@example.test','password':'Distinta!898765','rol':'ROLE','delegacion':'DELEG','delegaciones':'DELEGS'},'first_name',150,['username','first_name','email','password','rol','delegacion','delegaciones']),
 'items-metas':({'codigo':'NUEVO','nombre':'Nuevo Ítem','unidad_medida':'Cantidad'},'nombre',100,['codigo','nombre']),
 'tipos-atencion':({'codigo':'NUEVO','nombre':'Nuevo Tipo'},'nombre',100,['codigo','nombre']),
 'subatenciones':({'codigo':'NUEVO','nombre':'Nueva Sub','tipo_atencion':'TYPE'},'nombre',100,['codigo','nombre','tipo_atencion'])}

def variant_test(resource,variant):
    def test(self):
        raw,field,limit,required=SPECIFIC[resource]
        payload={k:self.role.pk if v=='ROLE' else self.d.pk if v=='DELEG' else [self.d.pk] if v=='DELEGS' else self.tipo.pk if v=='TYPE' else v for k,v in raw.items()}
        url='/api/'+resource+'/'
        if variant=='limite_exacto':
            payload[field]='Z'*limit;response=self.send(url,payload);self.assertEqual(response.status_code,201,response.content)
        elif variant=='longitud_excedida':
            payload[field]='Z'*(limit+1);response=self.send(url,payload);self.assertEqual(response.status_code,400,response.content)
        elif variant=='duplicado':
            first=self.send(url,payload);self.assertEqual(first.status_code,201,first.content)
            self.assertIn(self.send(url,payload).status_code,[400,409])
        elif variant=='obligatorios':
            for name in required:
                with self.subTest(campo=name):
                    invalid={k:v for k,v in payload.items() if k!=name}
                    response=self.send(url,invalid);self.assertEqual(response.status_code,400,response.content)
        elif variant=='formato_invalido':
            payload[field]={'dato':True};response=self.send(url,payload);self.assertEqual(response.status_code,400,response.content)
        elif variant=='estado_invalido':
            payload['estado']='Activo';response=self.send(url,payload);self.assertEqual(response.status_code,400,response.content)
    test.__doc__=f'A1/A3 — {resource}: {variant}; no guardar datos inválidos.'
    return test
for resource in SPECIFIC:
    for variant in ['limite_exacto','longitud_excedida','duplicado','obligatorios','formato_invalido','estado_invalido']:
        setattr(ValidacionesPorMaestro,'test_'+resource.replace('-','_')+'_'+variant,variant_test(resource,variant))

from django.test import TransactionTestCase
from django.db import close_old_connections
from concurrent.futures import ThreadPoolExecutor

class ConcurrenciaLocal(TransactionTestCase):
    def setUp(self):
        self.d=Delegacion.objects.create(codigo='C-D',nombre='Delegación Concurrencia')
        self.user=Usuario.objects.create_superuser('concurrente','concurrente@example.test','ClaveConcurrencia!372',delegacion=self.d)
        self.tipo=TipoAtencion.objects.create(codigo='C-T',nombre='Tipo Concurrencia')
        self.sub=SubAtencion.objects.create(codigo='C-S',nombre='Sub Concurrencia',tipo_atencion=self.tipo)

    def parallel_post(self,path,payload):
        close_old_connections()
        try:
            client=Client();client.force_login(Usuario.objects.get(pk=self.user.pk))
            return client.post(path,payload,content_type='application/json').status_code
        finally:close_old_connections()

    def test_dos_creaciones_duplicadas_solo_una_persiste(self):
        with ThreadPoolExecutor(max_workers=2) as pool:
            results=list(pool.map(lambda _:self.parallel_post('/api/tipos-atencion/',{'codigo':'CONC','nombre':'Duplicado Concurrente'}),range(2)))
        self.assertEqual(results.count(201),1);self.assertTrue(all(s in [201,400,409] for s in results),results)
        self.assertEqual(TipoAtencion.objects.filter(codigo='CONC').count(),1)

    def test_dos_envios_misma_clave_no_duplican(self):
        payload={'solicitante':'Persona Concurrente','tipo_atencion':self.tipo.pk,'sub_atencion':self.sub.pk,'descripcion':'Prueba simultánea','clave_idempotencia':'CONC-UNICO'}
        with ThreadPoolExecutor(max_workers=2) as pool:
            results=list(pool.map(lambda _:self.parallel_post('/api/atenciones/',payload),range(2)))
        self.assertEqual(sorted(results),[200,201],results)
        self.assertEqual(Atencion.objects.filter(clave_idempotencia='CONC-UNICO').count(),1)
