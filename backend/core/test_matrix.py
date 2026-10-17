from django.test import TestCase
from django.utils import timezone
from .models import Delegacion, Usuario, TipoAtencion, SubAtencion, ItemMeta
class MatrizExcel(TestCase):
 def setUp(self):
  self.d=Delegacion.objects.create(codigo='MX',nombre='Matriz local');self.u=Usuario.objects.create_superuser('matriz','matriz@example.test','ClaveLocal!839',delegacion=self.d);self.client.force_login(self.u)
  self.t=TipoAtencion.objects.create(codigo='MX',nombre='Social');self.s=SubAtencion.objects.create(codigo='MX',nombre='Orientación',tipo_atencion=self.t);self.item=ItemMeta.objects.create(codigo='MX',nombre='Ítem matriz')
 def post(self,url,data):return self.client.post(url,data,content_type='application/json')
 def attention(self,**extra):return self.post('/api/atenciones/',dict(solicitante='Persona QA',tipo_atencion=self.t.pk,sub_atencion=self.s.pk,descripcion='Prueba matriz',**extra))
 def test_distingue_visita_programada_y_realizada(self):
  r=self.attention(requiere_visita=True,fecha_visita_programada='2026-10-10',fecha_visita_realizada='2026-10-12');self.assertEqual(r.status_code,201,r.content);self.assertEqual(str(r.json()['data']['fecha_visita_programada']),'2026-10-10');self.assertEqual(str(r.json()['data']['fecha_visita_realizada']),'2026-10-12')
 def test_tipo_gestion_incompatible_con_etapa_rechazado(self):
  r=self.attention(caso_social='MATRIZ-CASO',tipo_gestion='3RA GESTION');self.assertEqual(r.status_code,400,r.content)
 def test_reclamo_persistido_versionado_y_filtrado(self):
  r=self.post('/api/reclamos/',{'clase':'Reclamo','fecha':str(timezone.localdate()),'vecino':'Persona QA','contacto':'912345678','ubicacion':'Centro','situacion':'Luminaria sin servicio','medio_escrito':True,'informado_a':'Unidad municipal'});self.assertEqual(r.status_code,201,r.content);data=r.json()['data']
  self.assertEqual(self.client.get('/api/reclamos/?search=Luminaria').json()['count'],1)
  self.assertEqual(self.client.patch(f"/api/reclamos/{data['id']}/",{'_version':999,'situacion':'Otro'},content_type='application/json').status_code,409)
 def test_plantillas_publican_campos_sin_trabajadores(self):
  r=self.client.get('/api/planillas/');self.assertEqual(r.status_code,200,r.content);self.assertIn('T_OC',[x['codigo'] for x in r.json()['results']]);self.assertNotIn('Ant-Man',str(r.json()))
 def test_tubo_conserva_unidad_y_codigo_foto(self):
  r=self.post('/api/operacion/compromisos/',{'fecha_comprometida':str(timezone.localdate()),'descripcion':'Compromiso','solicitante':'Persona QA','sector':'Centro','responsable':self.u.pk,'area_apoyo':'Unidad','origen':'Interno','unidad':1,'codigo_verificacion_foto':'REF-QA-001'});self.assertEqual(r.status_code,201,r.content);self.assertEqual(str(r.json()['data']['unidad']),'1');self.assertEqual(r.json()['data']['codigo_verificacion_foto'],'REF-QA-001')
 def test_unidad_del_tubo_no_es_cantidad_arbitraria(self):
  r=self.post('/api/operacion/compromisos/',{'fecha_comprometida':str(timezone.localdate()),'descripcion':'Compromiso','solicitante':'Persona QA','sector':'Centro','responsable':self.u.pk,'area_apoyo':'Unidad','origen':'Interno','unidad':2});self.assertEqual(r.status_code,400,r.content)
 def test_reclamos_fuera_de_delegacion_no_visibles(self):
  from .models import Reclamo
  other=Delegacion.objects.create(codigo='MX2',nombre='Otra matriz');r=Reclamo.objects.create(delegacion=other,autor=self.u,clase='Reclamo',fecha=timezone.localdate(),vecino='QA',situacion='No visible')
  self.assertEqual(self.client.get(f'/api/reclamos/{r.pk}/').status_code,404)
 def test_social_vinculo_ajeno_rechazado(self):
  from .models import Compromiso
  other=Delegacion.objects.create(codigo='MX3',nombre='Otra matriz');c=Compromiso.objects.create(delegacion=other,autor=self.u,responsable=self.u,descripcion='Ajeno',solicitante='QA',sector='Otro',fecha_comprometida=timezone.localdate(),area_apoyo='QA',origen='Interno')
  self.assertEqual(self.attention(compromiso=c.pk).status_code,400)
 def test_social_exporta_las_tres_etapas_y_no_formulas(self):
  from io import BytesIO
  from openpyxl import load_workbook
  for kind in ['1RA GESTION','2DA GESTION','3RA GESTION']:
   r=self.attention(caso_social='MATRIZ-EXPORT',tipo_gestion=kind,gestion_realizada='=QA');self.assertEqual(r.status_code,201,r.content)
  response=self.client.get('/api/casos-sociales/?export=1');self.assertEqual(response.status_code,200)
  sheet=load_workbook(BytesIO(response.content),data_only=False).active;self.assertEqual(sheet.max_row,2);self.assertEqual(sheet.cell(2,10).value,'1RA GESTION');self.assertEqual(sheet.cell(2,12).data_type,'s')
 def test_reclamos_solo_lectura_no_crea(self):
  from .models import Rol
  role=Rol.objects.create(nombre='Consulta matriz',permisos=['reclamos.ver']);u=Usuario.objects.create_user('consulta_matriz',rol=role,delegacion=self.d);self.client.force_login(u)
  self.assertEqual(self.post('/api/reclamos/',{}).status_code,403);self.assertEqual(self.client.get('/api/reclamos/opciones/').status_code,200)
 def test_planilla_asignada_no_permite_suplantar_otra(self):
  from .models import Cargo,Rol
  cargo=Cargo.objects.create(codigo='MX',nombre='Función matriz',area='QA',planilla='T_OC');role=Rol.objects.create(nombre='Captura matriz',permisos=['actividades.ver','actividades.crear']);u=Usuario.objects.create_user('funcion_matriz',rol=role,delegacion=self.d,cargo=cargo);self.client.force_login(u)
  r=self.post('/api/operacion/actividades/',{'fecha':str(timezone.localdate()),'solicitud':'QA','accion':'QA','contacto':'QA','telefono':'912345678','sector':'Centro','item':self.item.pk,'plantilla':'GES_SEG'});self.assertEqual(r.status_code,400,r.content)
 def test_editar_conserva_planilla_aunque_el_cargo_cambie(self):
  from .models import Actividad,Cargo,Rol
  a=Actividad.objects.create(autor=self.u,delegacion=self.d,item=self.item,fecha=timezone.localdate(),solicitud='Anterior',accion='QA',contacto='QA',telefono='912345678',sector='Centro',plantilla='T_OC')
  cargo=Cargo.objects.create(codigo='EDIT-MX',nombre='Otra función matriz',area='QA',planilla='AP_ADM');role=Rol.objects.create(nombre='Edición matriz',permisos=['actividades.ver','actividades.editar']);self.u.is_superuser=False;self.u.rol=role;self.u.cargo=cargo;self.u.save()
  r=self.client.patch(f'/api/operacion/actividades/{a.pk}/',{'_version':a.version,'plantilla':'T_OC','accion':'Actualizada'},content_type='application/json');self.assertEqual(r.status_code,200,r.content);self.assertEqual(r.json()['data']['plantilla'],'T_OC')
 def test_crear_sin_plantilla_toma_funcion_y_no_vista(self):
  from .models import Cargo,Rol
  cargo=Cargo.objects.create(codigo='NEW-MX',nombre='Función captura matriz',area='QA',planilla='T_OC');role=Rol.objects.create(nombre='Captura vista matriz',permisos=['actividades.ver','actividades.crear']);self.u.is_superuser=False;self.u.rol=role;self.u.cargo=cargo;self.u.save()
  r=self.post('/api/operacion/actividades/',{'fecha':str(timezone.localdate()),'solicitud':'QA','accion':'QA','contacto':'QA','telefono':'912345678','sector':'Centro','item':self.item.pk});self.assertEqual(r.status_code,201,r.content);self.assertEqual(r.json()['data']['plantilla'],'T_OC')
 def test_seguridad_asocia_profesional_configurado_sin_importar_personas(self):
  r=self.post('/api/operacion/actividades/',{'fecha':str(timezone.localdate()),'solicitud':'QA','accion':'QA','contacto':'QA','telefono':'912345678','sector':'Centro','item':self.item.pk,'plantilla':'GES_SEG','profesional':self.u.pk});self.assertEqual(r.status_code,201,r.content);self.assertEqual(r.json()['data']['profesional'],self.u.pk)
 def test_g_soc_abre_planilla_social_no_actividad_generica(self):
  r=self.client.get('/api/planillas/');profile=next(x for x in r.json()['results'] if x['codigo']=='G_SOC');self.assertEqual(profile.get('ruta'),'/pages/casos-sociales.html')
