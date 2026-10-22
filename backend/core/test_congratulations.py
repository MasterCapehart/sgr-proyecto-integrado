import tempfile
from datetime import date
from django.test import TestCase,override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from .models import Delegacion,Rol,Usuario,Reclamo

class CongratulationsTests(TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.settings=override_settings(MEDIA_ROOT=self.tmp.name);self.settings.enable();self.addCleanup(self.settings.disable)
  self.d=Delegacion.objects.create(codigo='FEL',nombre='Delegación QA')
  self.role=Rol.objects.create(nombre='Delegado',permisos=['reclamos.ver','reclamos.crear','reclamos.editar','reclamos.ver_equipo','reclamos.editar_equipo','reclamos.exportar'])
  self.u=Usuario.objects.create_user('delegado_fel',rol=self.role,delegacion=self.d);self.client.force_login(self.u)
 def send(self,url,data,method='post'):return getattr(self.client,method)(url,data,content_type='application/json')
 def payload(self,**extra):return dict(clase='Felicitación',fecha=str(timezone.localdate()),vecino='Persona ficticia',contacto='912345678',situacion='Agradecimiento por trabajo comunitario',procedencia='Directa',destinatario='Delegación')|extra
 def create(self,**extra):
  r=self.send('/api/reclamos/',self.payload(**extra));self.assertEqual(r.status_code,201,r.content);return r.json()['data']
 def file(self,row,content=b'%PDF-1.4\nRespaldo ficticio',name='carta.pdf'):
  return self.client.post(f"/api/reclamos/{row['id']}/respaldos/",{'_version':row['version'],'archivo':SimpleUploadedFile(name,content,content_type='application/pdf')})
 def review(self,row,**extra):return self.send(f"/api/reclamos/{row['id']}/revision/",dict(_version=row['version'],decision='Validada',observacion='Procedencia y respaldo comprobados',**extra))
 def test_otros_perfiles_siguen_restringidos(self):
  for name,superuser in [('Funcionario',False),('Verificador',False)]:
   role=Rol.objects.create(nombre=name,permisos=self.role.permisos);self.u.rol=role;self.u.is_superuser=superuser;self.u.save()
   self.assertEqual(self.send('/api/reclamos/',self.payload()).status_code,403)
  self.assertEqual(Reclamo.objects.count(),0)
 def test_alta_pendiente_y_opciones_servidor(self):
  row=self.create();self.assertEqual(row['revision']['estado'],'Pendiente');self.assertFalse(row['revision']['potencialmente_puntuable']);self.assertEqual(row['revisiones'],[])
  self.assertTrue(self.client.get('/api/reclamos/opciones/').json()['puede_gestionar_felicitaciones'])
 def test_historico_no_se_aprueba_y_no_cambia_clase(self):
  old=Reclamo.objects.create(autor=self.u,delegacion=self.d,clase='Felicitación',fecha=timezone.localdate(),vecino='QA',situacion='Histórico')
  r=self.client.get(f'/api/reclamos/{old.pk}/');self.assertFalse(r.json()['data']['revision']['potencialmente_puntuable'])
  r=self.send(f'/api/reclamos/{old.pk}/',dict(_version=old.version,clase='Reclamo'),'patch');self.assertEqual(r.status_code,400)
 def test_no_se_inyecta_aprobacion(self):
  p=self.payload();p['estado_revision']='Validada';self.assertEqual(self.send('/api/reclamos/',p).status_code,400)
 def test_validar_exige_respaldo_real(self):
  row=self.create();self.assertEqual(self.review(row).status_code,400)
 def test_carga_valida_revision_separada_y_descarga(self):
  row=self.create();r=self.file(row);self.assertEqual(r.status_code,201,r.content);row=r.json()['data'];self.assertFalse(row['revision']['potencialmente_puntuable']);self.assertEqual(len(row['respaldos']),1)
  r=self.review(row);self.assertEqual(r.status_code,200,r.content);row=r.json()['data'];self.assertTrue(row['revision']['potencialmente_puntuable']);self.assertEqual(row['revision']['estado'],'Validada');self.assertNotIn('puntaje',row['revision']);self.assertEqual(len(row['revisiones']),1)
  r=self.client.get(row['respaldos'][0]['archivo_url']);self.assertEqual(r.status_code,200);self.assertEqual(r['Cache-Control'],'private, no-store');self.assertTrue(b''.join(r.streaming_content).startswith(b'%PDF-'))
 def test_personal_y_destino_funcionario_no_son_candidatos(self):
  for key,value in [('procedencia','Personal'),('destinatario','Funcionario')]:
   row=self.create(**{key:value});r=self.file(row);self.assertEqual(r.status_code,201,r.content);row=r.json()['data'];self.assertEqual(self.review(row).status_code,400)
 def test_rechazo_requiere_motivo_y_se_conserva(self):
  row=self.create();r=self.send(f"/api/reclamos/{row['id']}/revision/",dict(_version=row['version'],decision='Rechazada',observacion=''));self.assertEqual(r.status_code,400)
  r=self.send(f"/api/reclamos/{row['id']}/revision/",dict(_version=row['version'],decision='Rechazada',observacion='No existe respaldo formal'));self.assertEqual(r.status_code,200,r.content);self.assertEqual(r.json()['data']['revision']['estado'],'Rechazada')
 def test_edicion_o_nuevo_respaldo_requiere_revalidar(self):
  row=self.file(self.create()).json()['data'];row=self.review(row).json()['data'];old=row['version']
  r=self.send(f"/api/reclamos/{row['id']}/",dict(_version=old,situacion='Descripción corregida'),'patch');self.assertEqual(r.status_code,200,r.content);row=r.json()['data'];self.assertEqual(row['revision']['estado'],'Pendiente');self.assertEqual(len(row['revisiones']),1)
  self.assertEqual(self.review(dict(row,version=old)).status_code,409)
  row=self.review(row).json()['data'];row=self.file(row,b'%PDF-1.4\nNueva carta').json()['data'];self.assertEqual(row['revision']['estado'],'Pendiente');self.assertEqual(len(row['respaldos']),2);self.assertEqual(len(row['revisiones']),2)
 def test_archivo_falso_rechazado_y_sin_mutacion(self):
  row=self.create();self.assertEqual(self.file(row,b'archivo falso').status_code,400);self.assertEqual(self.client.get(f"/api/reclamos/{row['id']}/").json()['data']['version'],row['version'])
 def test_aislamiento_y_restriccion_edicion_de_felicitacion(self):
  row=self.file(self.create()).json()['data'];url=row['respaldos'][0]['archivo_url']
  other=Delegacion.objects.create(codigo='FEL2',nombre='Otra');self.u.delegacion=other;self.u.save();self.assertEqual(self.client.get(url).status_code,404);self.assertEqual(self.review(row).status_code,404)
  self.u.delegacion=self.d;self.u.rol=Rol.objects.create(nombre='Funcionario',permisos=self.role.permisos);self.u.is_superuser=False;self.u.save();self.assertEqual(self.file(row).status_code,403);self.assertEqual(self.review(row).status_code,403);self.assertEqual(self.send(f"/api/reclamos/{row['id']}/",dict(_version=row['version'],situacion='Cambio'),'patch').status_code,403)
 def test_reclamo_normal_se_conserva_para_administrador(self):
  self.u.rol=None;self.u.is_superuser=True;self.u.save();p=self.payload();p.pop('procedencia');p.pop('destinatario');p['clase']='Reclamo';r=self.send('/api/reclamos/',p);self.assertEqual(r.status_code,201,r.content)
 def test_exportacion_incluye_estado_y_no_formula(self):
  from io import BytesIO
  from openpyxl import load_workbook
  row=self.create(vecino='=Nombre ficticio');response=self.client.get('/api/reclamos/export/');self.assertEqual(response.status_code,200);sheet=load_workbook(BytesIO(response.content)).active
  self.assertEqual(sheet.cell(2,4).data_type,'s');self.assertEqual(sheet.cell(2,14).value,'Pendiente');self.assertEqual(sheet.cell(2,15).value,'No')
 def test_no_se_acepta_evidencia_ajena_como_respaldo_formal(self):
  row=self.create(medio_escrito=True);self.assertEqual(self.review(row).status_code,400)
 def test_reclamo_no_admite_endpoint_de_felicitacion(self):
  p=self.payload();p['clase']='Reclamo';row=self.send('/api/reclamos/',p).json()['data'];self.assertEqual(self.file(row).status_code,404);self.assertEqual(self.review(row).status_code,404)
 def test_estado_rol_inactivo_bloquea_usuario_normal(self):
  self.role.estado=False;self.role.save();self.assertEqual(self.send('/api/reclamos/',self.payload()).status_code,403)
 def test_descarga_no_autenticada_y_respaldo_de_otro_registro(self):
  row=self.file(self.create()).json()['data'];other=self.create();backing=row['respaldos'][0]['id'];self.assertEqual(self.client.get(f"/api/reclamos/{other['id']}/respaldos/{backing}/archivo/").status_code,404)
  self.client.logout();self.assertEqual(self.client.get(row['respaldos'][0]['archivo_url']).status_code,401)
 def test_respaldo_desaparecido_impide_validar(self):
  from .models import RespaldoFelicitacion
  row=self.file(self.create()).json()['data'];b=RespaldoFelicitacion.objects.get(pk=row['respaldos'][0]['id']);b.archivo.storage.delete(b.archivo.name)
  self.assertEqual(self.review(row).status_code,400);self.assertEqual(self.client.get(f"/api/reclamos/{row['id']}/").json()['data']['revision']['estado'],'Pendiente')


 def test_administrador_y_superusuario_gestionan_flujo_completo(self):
  for name,superuser in [('Administrador',False),(None,True)]:
   self.u.rol=Rol.objects.create(nombre=name,permisos=self.role.permisos) if name else None;self.u.is_superuser=superuser;self.u.save()
   self.assertTrue(self.client.get('/api/reclamos/opciones/').json()['puede_gestionar_felicitaciones'])
   row=self.create();row=self.file(row).json()['data'];r=self.review(row);self.assertEqual(r.status_code,200,r.content);row=r.json()['data'];self.assertTrue(row['revision']['potencialmente_puntuable'])
   r=self.send(f"/api/reclamos/{row['id']}/",dict(_version=row['version'],situacion='Corrección del administrador'),'patch');self.assertEqual(r.status_code,200,r.content);self.assertEqual(r.json()['data']['revision']['estado'],'Pendiente')
 def test_superusuario_inactivo_no_gestiona(self):
  self.u.is_superuser=True;self.u.is_active=False;self.u.save();self.assertIn(self.send('/api/reclamos/',self.payload()).status_code,[401,403])

from django.test import TransactionTestCase,Client
from django.db import close_old_connections
from concurrent.futures import ThreadPoolExecutor
class ConcurrentCongratulationsTests(TransactionTestCase):
 setUp=CongratulationsTests.setUp
 send=CongratulationsTests.send
 payload=CongratulationsTests.payload
 create=CongratulationsTests.create
 file=CongratulationsTests.file
 def test_dos_revisiones_misma_version_una_sola_se_guarda(self):
  from .models import RevisionFelicitacion
  row=self.file(self.create()).json()['data'];second=Usuario.objects.create_user('otro_delegado_fel',rol=self.role,delegacion=self.d);clients=[]
  for user in [self.u,second]:
   c=Client();c.force_login(user);clients.append(c)
  def submit(c):
   close_old_connections()
   try:return c.post(f"/api/reclamos/{row['id']}/revision/",dict(_version=row['version'],decision='Validada',observacion='Comprobación formal de prueba'),content_type='application/json').status_code
   finally:close_old_connections()
  with ThreadPoolExecutor(max_workers=2) as pool:statuses=list(pool.map(submit,clients))
  self.assertEqual(sorted(statuses),[200,409]);self.assertEqual(RevisionFelicitacion.objects.count(),1)
