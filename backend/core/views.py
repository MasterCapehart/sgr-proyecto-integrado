import json
import random
import datetime
import openpyxl
from django.db import models, transaction
from django.utils import timezone
from django.views import View
from django.contrib.auth import authenticate, login, logout
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from django.utils.decorators import method_decorator
from django.core.mail import send_mail

from .models import (
    Usuario, TipoAtencion, SubAtencion, Delegacion,
    ItemMeta, ConfiguracionMeta, Atencion, AuditoriaDerivacion, Rol
)
from .api import BaseCRUDView

# ==========================================
# 1. Autenticación y Recuperación
# ==========================================

@csrf_exempt
def login_view(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            username = data.get('username')
            password = data.get('password')
            
            user = authenticate(request, username=username, password=password)
            
            if user is not None:
                if user.is_active:
                    login(request, user)
                    return JsonResponse({'mensaje': 'Login exitoso', 'usuario': user.username}, status=200)
                else:
                    return JsonResponse({'error': 'Cuenta inactiva'}, status=403)
            else:
                return JsonResponse({'error': 'Credenciales inválidas'}, status=401)
        except json.JSONDecodeError:
            return JsonResponse({'error': 'Formato JSON inválido'}, status=400)
            
    return JsonResponse({'error': 'Método no permitido. Use POST.'}, status=405)


@csrf_exempt
def logout_view(request):
    if request.method == 'POST':
        logout(request)
        return JsonResponse({'mensaje': 'Logout exitoso'}, status=200)
    return JsonResponse({'error': 'Método no permitido. Use POST.'}, status=405)


@csrf_exempt
def solicitar_recuperacion(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            email = data.get('email')
            
            if not email:
                return JsonResponse({'error': 'El correo es obligatorio'}, status=400)
            
            usuario = Usuario.objects.filter(email=email).first()
            
            if usuario:
                codigo = str(random.randint(100000, 999999))
                usuario.codigo_recuperacion = codigo
                usuario.save()
                
                send_mail(
                    'Código de Recuperación - SGR',
                    f'Hola {usuario.username},\n\nTu código de recuperación es: {codigo}\n\nSi no solicitaste esto, ignora este correo.',
                    'soporte@sgr.cl',
                    [email],
                    fail_silently=False,
                )
            
            return JsonResponse({'mensaje': 'Si el correo está registrado, recibirás un código pronto.'}, status=200)
            
        except json.JSONDecodeError:
            return JsonResponse({'error': 'Formato JSON inválido'}, status=400)
            
    return JsonResponse({'error': 'Método no permitido. Use POST.'}, status=405)


# ==========================================
# 2. Vistas del Patrón Común de CRUD (T04)
# ==========================================

class TipoAtencionCRUDView(BaseCRUDView):
    model = TipoAtencion
    fields = ["codigo", "nombre", "descripcion", "estado"]
    search_fields = ["codigo", "nombre", "descripcion"]
    ordering = ["codigo"]

    def validate_payload(self, data, instance=None):
        errors = super().validate_payload(data, instance=instance)
        
        codigo = data.get("codigo")
        if codigo:
            qs_codigo = TipoAtencion.objects.filter(codigo__iexact=codigo.strip())
            if instance:
                qs_codigo = qs_codigo.exclude(pk=instance.pk)
            if qs_codigo.exists():
                errors.setdefault("codigo", []).append("Ya existe un tipo de atención con este código.")

        nombre = data.get("nombre")
        if nombre:
            qs_nombre = TipoAtencion.objects.filter(nombre__iexact=nombre.strip())
            if instance:
                qs_nombre = qs_nombre.exclude(pk=instance.pk)
            if qs_nombre.exists():
                errors.setdefault("nombre", []).append("Ya existe un tipo de atención con este nombre.")

        return errors

    def delete(self, request, pk=None, *args, **kwargs):
        try:
            instance = self.get_queryset().get(pk=pk)
        except self.model.DoesNotExist:
            from .api import api_error
            return api_error(code="NOT_FOUND", message=f"El tipo de atención con ID {pk} no existe.", status=404)

        if hasattr(instance, "subatenciones") and instance.subatenciones.filter(estado=True).exists():
            from .api import api_error
            return api_error(
                code="HAS_ACTIVE_DEPENDENCIES",
                message="No se puede desactivar este tipo de atención porque posee subatenciones activas asociadas. Desactive primero las subatenciones hijas.",
                status=400
            )

        return super().delete(request, pk=pk, *args, **kwargs)

    def toggle_estado(self, request, pk):
        try:
            instance = self.get_queryset().get(pk=pk)
        except self.model.DoesNotExist:
            from .api import api_error
            return api_error(code="NOT_FOUND", message=f"El tipo de atención con ID {pk} no existe.", status=404)

        if instance.estado:
            if hasattr(instance, "subatenciones") and instance.subatenciones.filter(estado=True).exists():
                from .api import api_error
                return api_error(
                    code="HAS_ACTIVE_DEPENDENCIES",
                    message="No se puede desactivar este tipo de atención porque posee subatenciones activas asociadas. Desactive primero las subatenciones vinculadas.",
                    status=400
                )

        return super().toggle_estado(request, pk)


class DelegacionCRUDView(BaseCRUDView):
    model = Delegacion
    fields = ["nombre", "direccion", "comuna", "estado"]
    search_fields = ["nombre", "direccion", "comuna"]
    ordering = ["nombre"]

class RolCRUDView(BaseCRUDView):
    """CRUD para la gestión de Roles y Permisos (T09)"""
    model = Rol
    fields = ["nombre", "descripcion", "estado"]
    search_fields = ["nombre", "descripcion"]
    ordering = ["nombre"]

class UsuarioCRUDView(BaseCRUDView):
    """CRUD para la gestión de Usuarios"""
    model = Usuario
    fields = ["username", "first_name", "last_name", "email", "is_active"]
    search_fields = ["username", "first_name", "last_name", "email"]
    ordering = ["username"]

    def get_queryset(self):
        # Optimiza la consulta trayendo datos de llave foránea de una vez
        return Usuario.objects.select_related("rol", "delegacion").all()

    def serialize_instance(self, instance):
        data = super().serialize_instance(instance)
        # Asegura compatibilidad con el JS que busca el campo 'estado'
        data["estado"] = getattr(instance, 'is_active', getattr(instance, 'estado', True))
        
        # Formatea los datos para que el JS los inyecte directo en la tabla
        data["nombre_apellidos"] = instance.get_full_name() or instance.username
        data["correo_institucional"] = getattr(instance, 'email', getattr(instance, 'correo', 'Sin correo'))
        data["rol_asignado"] = instance.rol.nombre if getattr(instance, 'rol', None) else "Sin rol"
        data["delegacion_activa"] = instance.delegacion.nombre if getattr(instance, 'delegacion', None) else "Sin delegación"
        return data

    def export(self, request):
        """Exporta usuarios a XLSX real (Tarea T12)."""
        qs = self.filter_queryset(self.get_queryset())
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Usuarios"
        
        headers = ["ID", "Username", "Nombre y Apellidos", "Correo Institucional", "Rol Asignado", "Delegación Activa", "Estado"]
        ws.append(headers)
        
        for item in qs:
            estado_str = "Activo" if getattr(item, 'is_active', getattr(item, 'estado', True)) else "Inactivo"
            ws.append([
                item.pk,
                item.username,
                item.get_full_name() or item.username,
                getattr(item, 'email', getattr(item, 'correo', '')),
                item.rol.nombre if getattr(item, 'rol', None) else "",
                item.delegacion.nombre if getattr(item, 'delegacion', None) else "",
                estado_str
            ])
            
        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="usuarios_export.xlsx"'
        wb.save(response)
        return response


class SubAtencionCRUDView(BaseCRUDView):
    model = SubAtencion
    fields = ["codigo", "nombre", "descripcion", "tipo_atencion", "estado"]
    search_fields = ["codigo", "nombre", "descripcion", "tipo_atencion__codigo", "tipo_atencion__nombre"]
    ordering = ["codigo"]

    def get_queryset(self):
        return SubAtencion.objects.select_related("tipo_atencion").all()

    def filter_queryset(self, qs):
        qs = super().filter_queryset(qs)
        tipo_id = self.request.GET.get("tipo_atencion") or self.request.GET.get("tipo_padre")
        if tipo_id:
            if str(tipo_id).isdigit():
                qs = qs.filter(tipo_atencion_id=int(tipo_id))
            else:
                cod_candidate = str(tipo_id).split("·")[0].strip()
                qs = qs.filter(models.Q(tipo_atencion__codigo__iexact=cod_candidate) | models.Q(tipo_atencion__nombre__iexact=str(tipo_id).strip()))
        return qs

    def serialize_instance(self, instance):
        data = super().serialize_instance(instance)
        if instance.tipo_atencion:
            data["tipo_atencion"] = instance.tipo_atencion.pk
            data["tipo_atencion_codigo"] = instance.tipo_atencion.codigo
            data["tipo_atencion_nombre"] = instance.tipo_atencion.nombre
            data["tipo_padre_display"] = f"{instance.tipo_atencion.codigo} · {instance.tipo_atencion.nombre}"
        else:
            data["tipo_atencion"] = None
            data["tipo_padre_display"] = "Sin asignar"
        return data

    def validate_payload(self, data, instance=None):
        errors = {}
        tipo_id = data.get("tipo_atencion") or data.get("tipo_padre")
        parent_obj = None

        if not tipo_id:
            errors.setdefault("tipo_atencion", []).append("Debe seleccionar un Tipo de Atención (catálogo padre).")
        else:
            if str(tipo_id).isdigit():
                parent_obj = TipoAtencion.objects.filter(pk=int(tipo_id)).first()
            else:
                cod_candidate = str(tipo_id).split("·")[0].strip()
                parent_obj = TipoAtencion.objects.filter(codigo__iexact=cod_candidate).first()

            if not parent_obj:
                errors.setdefault("tipo_atencion", []).append("El Tipo de Atención indicado no existe.")
            elif not parent_obj.estado:
                errors.setdefault("tipo_atencion", []).append("El Tipo de Atención seleccionado está inactivo. Debe estar activo para asociarle subatenciones.")

        if parent_obj:
            data["tipo_atencion"] = parent_obj

        codigo = data.get("codigo")
        if not codigo:
            errors.setdefault("codigo", []).append("El código es obligatorio.")
        else:
            qs_cod = SubAtencion.objects.filter(codigo__iexact=str(codigo).strip())
            if instance:
                qs_cod = qs_cod.exclude(pk=instance.pk)
            if qs_cod.exists():
                errors.setdefault("codigo", []).append("Ya existe una subatención con este código.")

        nombre = data.get("nombre")
        if not nombre:
            errors.setdefault("nombre", []).append("El nombre es obligatorio.")
        elif parent_obj:
            qs_nom = SubAtencion.objects.filter(
                tipo_atencion=parent_obj,
                nombre__iexact=str(nombre).strip()
            )
            if instance:
                qs_nom = qs_nom.exclude(pk=instance.pk)
            if qs_nom.exists():
                errors.setdefault("nombre", []).append(f"Ya existe una subatención con este nombre bajo el tipo {parent_obj.codigo}.")

        return errors

    def export(self, request):
        """Exporta subatenciones a XLSX real (Tarea T12)."""
        qs = self.filter_queryset(self.get_queryset())
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Subatenciones"
        
        headers = ["ID", "Código", "Nombre de la Subatención", "Tipo de Atención Padre", "Estado"]
        ws.append(headers)
        
        for item in qs:
            padre_str = f"{item.tipo_atencion.codigo} · {item.tipo_atencion.nombre}" if item.tipo_atencion else ""
            ws.append([
                item.pk,
                item.codigo,
                item.nombre,
                padre_str,
                "Activo" if item.estado else "Inactivo"
            ])
            
        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="subatenciones_export.xlsx"'
        wb.save(response)
        return response


class ItemMetaCRUDView(BaseCRUDView):
    model = ItemMeta
    fields = ["codigo", "nombre", "descripcion", "unidad_medida", "estado"]
    search_fields = ["codigo", "nombre", "descripcion"]
    ordering = ["codigo"]

    def validate_payload(self, data, instance=None):
        errors = super().validate_payload(data, instance=instance)
        codigo = data.get("codigo")
        if codigo:
            qs_codigo = ItemMeta.objects.filter(codigo__iexact=str(codigo).strip())
            if instance:
                qs_codigo = qs_codigo.exclude(pk=instance.pk)
            if qs_codigo.exists():
                errors.setdefault("codigo", []).append("Ya existe un ítem con este código.")

        nombre = data.get("nombre")
        if nombre:
            qs_nombre = ItemMeta.objects.filter(nombre__iexact=str(nombre).strip())
            if instance:
                qs_nombre = qs_nombre.exclude(pk=instance.pk)
            if qs_nombre.exists():
                errors.setdefault("nombre", []).append("Ya existe un ítem con este nombre.")

        return errors

    def delete(self, request, pk=None, *args, **kwargs):
        try:
            instance = self.get_queryset().get(pk=pk)
        except self.model.DoesNotExist:
            from .api import api_error
            return api_error(code="NOT_FOUND", message=f"El ítem con ID {pk} no existe.", status=404)

        if hasattr(instance, "configuraciones_meta") and instance.configuraciones_meta.filter(estado=True).exists():
            from .api import api_error
            return api_error(
                code="HAS_ACTIVE_DEPENDENCIES",
                message="No se puede desactivar este ítem porque posee metas activas asignadas. Desactive primero las configuraciones de meta asociadas.",
                status=400
            )

        return super().delete(request, pk=pk, *args, **kwargs)

    def toggle_estado(self, request, pk):
        try:
            instance = self.get_queryset().get(pk=pk)
        except self.model.DoesNotExist:
            from .api import api_error
            return api_error(code="NOT_FOUND", message=f"El ítem con ID {pk} no existe.", status=404)

        if instance.estado:
            if hasattr(instance, "configuraciones_meta") and instance.configuraciones_meta.filter(estado=True).exists():
                from .api import api_error
                return api_error(
                    code="HAS_ACTIVE_DEPENDENCIES",
                    message="No se puede desactivar este ítem porque posee metas activas asignadas. Desactive primero las configuraciones de meta asociadas.",
                    status=400
                )

        return super().toggle_estado(request, pk)

    def export(self, request):
        """Exporta ítems y metas a XLSX real (Tarea T12)."""
        qs = self.filter_queryset(self.get_queryset())
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Items Metas"
        
        headers = ["ID", "Código", "Nombre del Ítem", "Descripción Operativa", "Unidad de Medida", "Estado"]
        ws.append(headers)
        
        for item in qs:
            ws.append([
                item.pk,
                item.codigo,
                item.nombre,
                item.descripcion or "",
                item.unidad_medida,
                "Activo" if item.estado else "Inactivo"
            ])
            
        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="items_metas_export.xlsx"'
        wb.save(response)
        return response


# ==========================================
# 5. Registro y Operación de Atención Ciudadana (Tarea T17)
# ==========================================

class AtencionContextoView(View):
    def get(self, request, *args, **kwargs):
        user = getattr(request, 'user', None)
        if user and not getattr(user, 'is_authenticated', False):
            user = None
        if not user:
            user = Usuario.objects.filter(is_active=True).first()

        del_sesion_id = request.session.get('delegacion_activa_id') if hasattr(request, 'session') else None
        delegacion = None
        if del_sesion_id:
            delegacion = Delegacion.objects.filter(pk=del_sesion_id, estado=True).first()

        if not delegacion:
            delegacion = getattr(user, 'delegacion', None) if user else None
        if not delegacion or not delegacion.estado:
            delegacion = Delegacion.objects.filter(estado=True).first()

        ahora = timezone.now()
        user_display = (
            f"{user.get_full_name() or user.username} ({user.rol.nombre if user.rol else 'Funcionario Municipal'})"
            if user else "Funcionario Municipal"
        )
        delegacion_display = delegacion.nombre if delegacion else "Delegación Las Compañías"

        return JsonResponse({
            "success": True,
            "data": {
                "usuario": {
                    "id": user.id if user else None,
                    "username": user.username if user else "operador",
                    "nombre_completo": user.get_full_name() if (user and user.get_full_name()) else (user.username if user else "Funcionario"),
                    "rol": user.rol.nombre if (user and user.rol) else "Operador Territorial",
                    "display": user_display
                },
                "delegacion": {
                    "id": delegacion.id if delegacion else None,
                    "nombre": delegacion.nombre if delegacion else "Delegación Las Compañías",
                    "display": delegacion_display
                },
                "fecha_servidor": ahora.isoformat(),
                "fecha_display": ahora.strftime("%d/%m/%Y %H:%M"),
                "folio_estimado": "Asignado automáticamente al guardar (AT-YYYY-XXXX)"
            }
        }, status=200)


class AtencionCRUDView(BaseCRUDView):
    model = Atencion
    fields = ["solicitante", "descripcion", "caso_social"]
    search_fields = ["folio", "solicitante", "descripcion", "tipo_atencion__nombre", "sub_atencion__nombre"]
    ordering = ["-id"]

    def get_queryset(self):
        return Atencion.objects.select_related("tipo_atencion", "sub_atencion", "delegacion", "usuario").all()

    def filter_queryset(self, qs):
        qs = super().filter_queryset(qs)
        del_param = self.request.GET.get("delegacion")
        if del_param:
            qs = qs.filter(delegacion_id=del_param)
        elif self.request.GET.get("solo_activa") == "1" or self.request.GET.get("ambito_activo") == "true":
            activa_id = self.request.session.get("delegacion_activa_id") if hasattr(self.request, "session") else None
            if activa_id:
                qs = qs.filter(delegacion_id=activa_id)
        return qs

    def serialize_instance(self, instance):
        derivaciones = [
            {
                "id": d.pk,
                "origen": d.delegacion_origen.nombre,
                "destino": d.delegacion_destino.nombre,
                "responsable": d.usuario_responsable.get_full_name() or d.usuario_responsable.username,
                "motivo": d.motivo,
                "fecha": d.fecha_derivacion.isoformat()
            }
            for d in getattr(instance, 'derivaciones_auditadas', []).all()
        ] if hasattr(instance, 'derivaciones_auditadas') else []

        return {
            "id": instance.pk,
            "folio": instance.folio,
            "codigo_atencion": instance.folio,
            "solicitante": instance.solicitante,
            "tipo_atencion": {
                "id": instance.tipo_atencion.pk,
                "codigo": instance.tipo_atencion.codigo,
                "nombre": instance.tipo_atencion.nombre,
                "display": f"{instance.tipo_atencion.codigo} · {instance.tipo_atencion.nombre}"
            },
            "sub_atencion": {
                "id": instance.sub_atencion.pk,
                "codigo": instance.sub_atencion.codigo,
                "nombre": instance.sub_atencion.nombre,
                "display": f"{instance.sub_atencion.codigo} · {instance.sub_atencion.nombre}"
            },
            "descripcion": instance.descripcion,
            "delegacion": {
                "id": instance.delegacion.pk,
                "nombre": instance.delegacion.nombre
            },
            "delegacion_responsable": {
                "id": instance.delegacion_responsable.pk,
                "nombre": instance.delegacion_responsable.nombre
            } if instance.delegacion_responsable else None,
            "usuario": {
                "id": instance.usuario.pk,
                "nombre": instance.usuario.get_full_name() or instance.usuario.username
            },
            "caso_social": instance.caso_social or "",
            "numero_gestion": instance.numero_gestion,
            "clave_idempotencia": instance.clave_idempotencia or "",
            "derivaciones_auditadas": derivaciones,
            "fecha_registro": instance.fecha_registro.isoformat() if instance.fecha_registro else None,
            "estado": instance.estado
        }

    def _generar_folio(self):
        anio = timezone.now().year
        prefijo = f"AT-{anio}-"
        ultimo = Atencion.objects.filter(folio__startswith=prefijo).order_by("-id").first()
        if ultimo and ultimo.folio:
            try:
                num = int(ultimo.folio.split("-")[-1]) + 1
            except (ValueError, IndexError):
                num = Atencion.objects.count() + 1
        else:
            num = 1
        return f"{prefijo}{num:04d}"

    def post(self, request, pk=None, action=None, *args, **kwargs):
        if action == "toggle-estado" and pk is not None:
            return self.toggle_estado(request, pk)

        from .api import api_error, api_success
        data = request.json_data
        errors = {}

        clave_idempotencia = (
            data.get("clave_idempotencia") or
            request.headers.get("X-Idempotency-Key") or
            request.META.get("HTTP_X_IDEMPOTENCY_KEY")
        )
        if clave_idempotencia:
            clave_idempotencia = str(clave_idempotencia).strip()
            existente = Atencion.objects.filter(clave_idempotencia=clave_idempotencia).first()
            if existente:
                return api_success(
                    data=self.serialize_instance(existente),
                    message=f"Atención procesada previamente (folio {existente.folio}). Se evitó el duplicado por reenvío.",
                    extra={"is_duplicate_prevented": True, "idempotent": True},
                    status=200
                )

        solicitante = str(data.get("solicitante") or "").strip()
        if not solicitante:
            errors.setdefault("solicitante", []).append("El nombre o RUT del solicitante es obligatorio.")

        descripcion = str(data.get("descripcion") or "").strip()
        if not descripcion:
            errors.setdefault("descripcion", []).append("La descripción de la atención es obligatoria.")

        tipo_id = data.get("tipo_atencion") or data.get("tipo")
        tipo_obj = None
        if not tipo_id:
            errors.setdefault("tipo_atencion", []).append("Debe seleccionar un Tipo de Atención.")
        else:
            if str(tipo_id).isdigit():
                tipo_obj = TipoAtencion.objects.filter(pk=int(tipo_id)).first()
            else:
                cod_candidate = str(tipo_id).split("·")[0].strip()
                tipo_obj = TipoAtencion.objects.filter(
                    models.Q(codigo__iexact=cod_candidate) | models.Q(nombre__iexact=str(tipo_id).strip())
                ).first()

            if not tipo_obj:
                errors.setdefault("tipo_atencion", []).append("El Tipo de Atención indicado no existe.")
            elif not tipo_obj.estado:
                errors.setdefault("tipo_atencion", []).append("El Tipo de Atención seleccionado está inactivo.")

        sub_id = data.get("sub_atencion") or data.get("subtipo")
        sub_obj = None
        if not sub_id:
            errors.setdefault("sub_atencion", []).append("Debe seleccionar una Subatención específica.")
        else:
            if str(sub_id).isdigit():
                sub_obj = SubAtencion.objects.filter(pk=int(sub_id)).first()
            else:
                cod_candidate = str(sub_id).split("·")[0].strip()
                sub_obj = SubAtencion.objects.filter(
                    models.Q(codigo__iexact=cod_candidate) | models.Q(nombre__iexact=str(sub_id).strip())
                ).first()

            if not sub_obj:
                errors.setdefault("sub_atencion", []).append("La Subatención indicada no existe.")
            elif not sub_obj.estado:
                errors.setdefault("sub_atencion", []).append("La Subatención seleccionada está inactiva.")
            elif tipo_obj and sub_obj.tipo_atencion_id != tipo_obj.pk:
                errors.setdefault("sub_atencion", []).append(
                    f"La subatención '{sub_obj.nombre}' no pertenece al tipo de atención '{tipo_obj.nombre}'."
                )

        if errors:
            return api_error(code="VALIDATION_ERROR", message="Error de validación.", details=errors, status=400)

        usuario = getattr(request, 'user', None)
        if usuario and not getattr(usuario, 'is_authenticated', False):
            usuario = None
        if not usuario:
            usuario = Usuario.objects.filter(is_active=True).first()

        if not usuario or not usuario.is_active:
            return api_error(code="RESPONSABLE_NO_AUTORIZADO", message="Usuario no activo o autorizado.", status=403)

        if usuario.rol and not usuario.rol.estado:
            return api_error(code="RESPONSABLE_NO_AUTORIZADO", message="Rol del usuario inactivo.", status=403)

        del_sesion_id = request.session.get('delegacion_activa_id') if hasattr(request, 'session') else None
        delegacion_origen = None
        if del_sesion_id:
            delegacion_origen = Delegacion.objects.filter(pk=del_sesion_id, estado=True).first()

        if not delegacion_origen:
            delegacion_origen = getattr(usuario, "delegacion", None)
        if not delegacion_origen or not delegacion_origen.estado:
            delegacion_origen = Delegacion.objects.filter(estado=True).first()

        if not delegacion_origen:
            return api_error(code="SERVER_CONTEXT_ERROR", message="Ámbito territorial no determinado.", status=500)

        del_manipulada = data.get("delegacion") or data.get("delegacion_origen") or data.get("id_delegacion")
        if del_manipulada:
            del_manipulada_id = None
            if str(del_manipulada).isdigit():
                del_manipulada_id = int(del_manipulada)
            else:
                d_found = Delegacion.objects.filter(models.Q(nombre__icontains=str(del_manipulada).strip())).first()
                if d_found:
                    del_manipulada_id = d_found.pk

            if del_manipulada_id and del_manipulada_id != delegacion_origen.pk:
                return api_error(
                    code="AMBITO_MANIPULADO",
                    message="No autorizado para registrar fuera de su delegación de origen.",
                    status=403
                )

        del_resp_val = data.get("delegacion_responsable")
        delegacion_responsable = delegacion_origen
        if del_resp_val:
            if str(del_resp_val).isdigit():
                delegacion_responsable = Delegacion.objects.filter(pk=int(del_resp_val)).first()
            else:
                delegacion_responsable = Delegacion.objects.filter(models.Q(nombre__icontains=str(del_resp_val).strip())).first()

            if not delegacion_responsable:
                return api_error(code="DELEGACION_RESPONSABLE_INVALIDA", message="Delegación responsable no existe.", status=400)
            if not delegacion_responsable.estado:
                return api_error(code="DELEGACION_RESPONSABLE_INACTIVA", message="Delegación responsable inactiva.", status=400)

        if not clave_idempotencia:
            ventana = timezone.now() - datetime.timedelta(seconds=4)
            duplicado_reciente = Atencion.objects.filter(
                solicitante__iexact=solicitante,
                tipo_atencion=tipo_obj,
                sub_atencion=sub_obj,
                descripcion__iexact=descripcion,
                usuario=usuario,
                fecha_registro__gte=ventana
            ).first()
            if duplicado_reciente:
                return api_success(
                    data=self.serialize_instance(duplicado_reciente),
                    message=f"Envío repetido detectado (folio {duplicado_reciente.folio}). Se evitó el duplicado.",
                    extra={"is_duplicate_prevented": True, "idempotent": True},
                    status=200
                )

        caso_social_raw = str(data.get("caso_social") or "").strip()
        caso_social_code = caso_social_raw.upper() if caso_social_raw else None

        with transaction.atomic():
            numero_gestion = None
            if caso_social_code:
                gestiones_anteriores = Atencion.objects.select_for_update().filter(caso_social__iexact=caso_social_code)
                conteo_gestiones = gestiones_anteriores.count()

                if conteo_gestiones >= 3:
                    return api_error(
                        code="CUARTA_GESTION_RECHAZADA",
                        message=f"El caso social '{caso_social_code}' ya cuenta con 3 gestiones. Límite máximo.",
                        status=400
                    )

                numero_gestion = conteo_gestiones + 1

            folio_real = self._generar_folio()

            atencion = Atencion.objects.create(
                folio=folio_real,
                solicitante=solicitante,
                tipo_atencion=tipo_obj,
                sub_atencion=sub_obj,
                descripcion=descripcion,
                delegacion=delegacion_origen,
                delegacion_responsable=delegacion_responsable,
                usuario=usuario,
                caso_social=caso_social_code,
                numero_gestion=numero_gestion,
                clave_idempotencia=clave_idempotencia or None,
                estado=True
            )

            if delegacion_responsable and delegacion_responsable.pk != delegacion_origen.pk:
                AuditoriaDerivacion.objects.create(
                    atencion=atencion,
                    delegacion_origen=delegacion_origen,
                    delegacion_destino=delegacion_responsable,
                    usuario_responsable=usuario,
                    motivo=str(data.get("motivo_derivacion") or "Derivación institucional autorizada").strip()
                )

        return api_success(data=self.serialize_instance(atencion), message=f"Atención registrada: {atencion.folio}", status=201)

    def export(self, request):
        """Exporta atenciones ciudadanas a XLSX real (Tarea T12)."""
        qs = self.filter_queryset(self.get_queryset())
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Atenciones"
        
        headers = ["ID", "Folio", "Solicitante", "Tipo de Atención", "Subatención", "Delegación", "Funcionario", "Fecha Registro", "Caso Social", "Estado"]
        ws.append(headers)
        
        for a in qs:
            ws.append([
                a.pk,
                a.folio,
                a.solicitante,
                a.tipo_atencion.nombre if a.tipo_atencion else "",
                a.sub_atencion.nombre if a.sub_atencion else "",
                a.delegacion.nombre if a.delegacion else "",
                a.usuario.get_full_name() or a.usuario.username if a.usuario else "",
                a.fecha_registro.strftime("%Y-%m-%d %H:%M:%S") if a.fecha_registro else "",
                a.caso_social or "",
                "Activo" if a.estado else "Inactivo"
            ])
            
        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="atenciones_export.xlsx"'
        wb.save(response)
        return response


# ==========================================
# 6. Contexto Institucional y Ámbito Territorial (Tarea T19)
# ==========================================

class ContextoView(View):
    def get(self, request, *args, **kwargs):
        user = getattr(request, 'user', None)
        if user and not getattr(user, 'is_authenticated', False):
            user = None
        if not user:
            user = Usuario.objects.filter(is_active=True).first()

        del_sesion_id = request.session.get('delegacion_activa_id') if hasattr(request, 'session') else None
        delegacion = None
        if del_sesion_id:
            delegacion = Delegacion.objects.filter(pk=del_sesion_id, estado=True).first()

        if not delegacion:
            delegacion = getattr(user, 'delegacion', None) if user else None
        if not delegacion or not delegacion.estado:
            delegacion = Delegacion.objects.filter(estado=True).first()

        if delegacion and hasattr(request, 'session'):
            request.session['delegacion_activa_id'] = delegacion.pk

        delegaciones_qs = Delegacion.objects.filter(estado=True).order_by('id')
        delegaciones_autorizadas = [
            {
                "id": d.pk,
                "nombre": d.nombre,
                "direccion": d.direccion or "",
                "comuna": d.comuna or "",
                "activa": (d.pk == delegacion.pk) if delegacion else False
            }
            for d in delegaciones_qs
        ]

        ahora = timezone.now()
        user_display = (
            f"{user.get_full_name() or user.username} ({user.rol.nombre if user.rol else 'Funcionario Municipal'})"
            if user else "Funcionario Municipal"
        )

        return JsonResponse({
            "success": True,
            "data": {
                "usuario": {
                    "id": user.pk if user else None,
                    "username": user.username if user else "operador",
                    "nombre_completo": user.get_full_name() if (user and user.get_full_name()) else (user.username if user else "Funcionario"),
                    "rut": getattr(user, 'rut', '') or '12.345.678-9',
                    "rol": user.rol.nombre if (user and user.rol) else "Operador Territorial",
                    "display": user_display
                },
                "delegacion_activa": {
                    "id": delegacion.pk if delegacion else None,
                    "nombre": delegacion.nombre if delegacion else "Delegación Las Compañías",
                    "display": delegacion.nombre if delegacion else "DEL01 · Las Compañías"
                },
                "delegaciones_autorizadas": delegaciones_autorizadas,
                "fecha_servidor": ahora.isoformat(),
                "fecha_display": ahora.strftime("%d/%m/%Y %H:%M")
            }
        }, status=200)


@method_decorator(csrf_exempt, name='dispatch')
class CambiarDelegacionView(View):
    def post(self, request, *args, **kwargs):
        from .api import api_error, api_success

        try:
            body = request.body.decode('utf-8')
            data = json.loads(body) if body else {}
        except Exception:
            return api_error(code="INVALID_JSON", message="Cuerpo JSON inválido.", status=400)

        del_id = data.get("delegacion_id")
        if not del_id:
            return api_error(code="DELEGACION_REQUIRED", message="ID de delegación requerido.", status=400)

        delegacion = Delegacion.objects.filter(pk=del_id).first()
        if not delegacion:
            return api_error(code="DELEGACION_NOT_FOUND", message="Delegación no existe.", status=404)

        if not delegacion.estado:
            return api_error(code="DELEGACION_INACTIVA", message="Delegación inactiva.", status=400)

        if hasattr(request, 'session'):
            request.session['delegacion_activa_id'] = delegacion.pk

        user = getattr(request, 'user', None)
        if user and getattr(user, 'is_authenticated', False):
            user.delegacion = delegacion
            user.save(update_fields=['delegacion'])

        return api_success(
            data={"id": delegacion.pk, "nombre": delegacion.nombre, "display": delegacion.nombre},
            message=f"Ámbito territorial cambiado a {delegacion.nombre}.",
            status=200
        )