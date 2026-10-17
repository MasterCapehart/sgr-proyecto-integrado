"""Seguimiento de derivaciones y expedientes sociales del contexto autorizado."""
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Q, Count, Max
from django.views.decorators.http import require_http_methods
from django.utils import timezone
from . import local_api as api
from . import domain_api as domain
from .models import Actividad, DerivacionActividad, Atencion

@require_http_methods(['POST'])
def derivation(request,pk):
    denied=domain.gate(request,'actividades','editar')
    if denied:return denied
    try:
        data=api.body(request)
        if set(data)-{'_version','destino','motivo','estado'}:return api.error('Campos de derivación no permitidos.')
        with transaction.atomic():
            obj=domain.scoped(request,'actividades').select_for_update().filter(pk=pk).first()
            if not obj:return api.error('Actividad no encontrada.',404)
            if obj.autor_id!=request.user.pk and not api.allowed(request.user,'actividades','editar_equipo'):return api.error('No puedes derivar una actividad ajena.',403)
            conflict=domain.validate_version(obj,data)
            if conflict:return conflict
            for key,limit in [('destino',200),('motivo',2000)]:
                if not isinstance(data.get(key),str) or not data[key].strip() or len(data[key])>limit:return api.error(f'{key}: texto obligatorio de hasta {limit} caracteres.')
            if data.get('estado') not in ['Pendiente','En proceso','Resuelta']:return api.error('Estado de derivación inválido.')
            previous=domain.representation(obj)
            event=DerivacionActividad(actividad=obj,autor=request.user,destino=data['destino'].strip(),motivo=data['motivo'].strip(),estado=data['estado']);event.full_clean();event.save()
            obj.derivado_a=event.destino;obj.estado_derivacion=event.estado;obj.version+=1;obj.save(update_fields=['derivado_a','estado_derivacion','version'])
            api.record_event(request,'actividades',obj.pk,'derivar',previous,domain.representation(obj))
        return api.ok(domain.representation(obj))
    except (ValidationError,ValueError,TypeError):return api.error('Datos de derivación inválidos.')

@require_http_methods(['GET'])
def cases(request,code=None):
    denied=domain.gate(request,'atenciones','ver')
    if denied:return denied
    d=api.active_delegation(request)
    qs=Atencion.objects.filter(Q(delegacion=d)|Q(delegacion_responsable=d)).exclude(caso_social__isnull=True).exclude(caso_social='')
    if code:
        records=qs.filter(caso_social=code.upper()).order_by('numero_gestion','id')
        first=records.first()
        if first is None:return api.error('Expediente no encontrado.',404)
        return api.ok({'codigo':first.caso_social,'solicitante':first.solicitante,'rut_solicitante':records.exclude(rut_solicitante='').values_list('rut_solicitante',flat=True).first() or '', 'telefono_solicitante':records.exclude(telefono_solicitante='').values_list('telefono_solicitante',flat=True).last() or '','gestiones':[api.serialize(o,api.RESOURCES['atenciones'][1]) for o in records]})
    search=request.GET.get('search','').strip()
    if search:
        codes=qs.filter(Q(caso_social__icontains=search)|Q(solicitante__icontains=search)|Q(rut_solicitante__icontains=search)).values('caso_social')
        qs=qs.filter(caso_social__in=codes)
    if request.GET.get('export')=='1':return social_export(request,qs)
    grouped=qs.values('caso_social').annotate(solicitante=Max('solicitante'),total_gestiones=Count('id'),ultima_fecha=Max('fecha_registro')).order_by('-ultima_fecha','caso_social')
    try:page=int(request.GET.get('page',1))
    except (ValueError,TypeError):return api.error('Página inválida.')
    pager=Paginator(grouped,20)
    if page<1 or page>pager.num_pages:return api.error('Página inexistente.',404)
    results=[]
    for row in pager.page(page):
        row.update(case_summary(qs.filter(caso_social=row['caso_social']).order_by('numero_gestion','id')));results.append(row)
    return api.ok(results=results,count=pager.count,page=page,total_pages=pager.num_pages,columnas=SOCIAL_COLUMNS)

SOCIAL_COLUMNS=[('caso_social','Código caso'),('fecha_ingreso','Fecha ingreso'),('tipo_nombre','Tipo atención'),('subatencion_nombre','Sub atención'),('solicitante','Usuario / vecino'),('rut','RUT'),('telefono','Teléfono'),('requiere_visita','Requiere visita'),('compromiso','Ingreso al tubo'),('primera','Primera gestión'),('programada','Fecha programada visita'),('observaciones','Observaciones'),('segunda','Segunda gestión'),('realizada','Fecha visita realizada'),('informe','Fecha entrega informe'),('tercera','Tercera gestión'),('beneficio','Fecha entrega beneficio')]
def case_summary(records):
    rows=list(records);first=rows[0]
    def latest(name):return next((getattr(r,name) for r in reversed(rows) if getattr(r,name) not in ['',None]),None)
    from .local_api import serialize,RESOURCES
    initial=serialize(first,RESOURCES['atenciones'][1])
    result={'caso_social':first.caso_social,'fecha_ingreso':timezone.localtime(first.fecha_registro).date().isoformat(),'tipo_nombre':initial['tipo_nombre'],'subatencion_nombre':initial['subatencion_nombre'],'solicitante':first.solicitante,'rut':latest('rut_solicitante') or '', 'telefono':latest('telefono_solicitante') or '', 'requiere_visita':'No informado' if latest('requiere_visita') is None else 'Sí' if latest('requiere_visita') else 'No','compromiso':latest('compromiso_id') or '', 'programada':latest('fecha_visita_programada'),'realizada':latest('fecha_visita_realizada'),'informe':latest('fecha_entrega_informe'),'beneficio':latest('fecha_entrega_beneficio'),'observaciones':latest('gestion_realizada') or first.descripcion}
    for number,key in [(1,'primera'),(2,'segunda'),(3,'tercera')]:
        row=next((r for r in rows if r.numero_gestion==number),None);result[key]=(row.tipo_gestion or 'Registrada sin tipo especificado') if row else 'Sin registrar'
    return result

def social_export(request,qs):
    from io import BytesIO
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    if not api.allowed(request.user,'atenciones','exportar'):return api.error('No tienes permiso para exportar atenciones.',403)
    book=Workbook();sheet=book.active;sheet.title='Gestión social';sheet.append([label for key,label in SOCIAL_COLUMNS]);codes=list(qs.order_by('caso_social').values_list('caso_social',flat=True).distinct())
    for code in codes:
        row=case_summary(qs.filter(caso_social=code).order_by('numero_gestion','id'));sheet.append([row[key] for key,label in SOCIAL_COLUMNS])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='AD0000')
    for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=27
    sheet.freeze_panes='F2';sheet.auto_filter.ref=sheet.dimensions;out=BytesIO();book.save(out);response=HttpResponse(out.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet');response['Content-Disposition']='attachment; filename="gestion-social.xlsx"';api.record_event(request,'atenciones',None,'exportar_social',None,{'casos':len(codes)});return response
